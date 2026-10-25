# -*- coding: utf-8 -*-

import base64
import io
from collections import OrderedDict

from dateutil.relativedelta import relativedelta

from odoo import api, fields, models
from odoo.exceptions import UserError


class PhanHeItCostCategory(models.Model):
    _name = "phan.he.it.cost.category"
    _description = "Hạng mục chi phí IT"
    _order = "sequence, name"

    name = fields.Char(string="Hạng mục", required=True)
    code = fields.Char(string="Mã", required=True)
    sequence = fields.Integer(default=10)
    active = fields.Boolean(default=True)

    _code_uniq = models.Constraint("unique(code)", "Mã hạng mục phải duy nhất.")


class PhanHeItRenewalCycle(models.Model):
    _name = "phan.he.it.renewal.cycle"
    _description = "Chu kỳ gia hạn"
    _order = "months, name"

    name = fields.Char(string="Tên chu kỳ", required=True)
    code = fields.Char(string="Mã", required=True)
    months = fields.Integer(string="Số tháng", required=True, default=12)
    active = fields.Boolean(default=True)

    _code_uniq = models.Constraint("unique(code)", "Mã chu kỳ phải duy nhất.")


class PhanHeItCost(models.Model):
    _name = "phan.he.it.cost"
    _description = "Khoản chi phí / gia hạn tài sản IT"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "date_end asc, code, id desc"
    _rec_names_search = ["code", "contract_ref"]

    code = fields.Char(
        string="Mã khoản chi phí",
        required=True,
        copy=False,
        default=lambda self: self.env["ir.sequence"].next_by_code("phan.he.it.cost") or "New",
        tracking=True,
    )
    asset_id = fields.Many2one(
        "phan.he.it.asset",
        string="Tài sản",
        required=True,
        ondelete="restrict",
        index=True,
        tracking=True,
    )
    category_id = fields.Many2one(
        "phan.he.it.cost.category",
        string="Hạng mục chi phí",
        required=True,
        ondelete="restrict",
        tracking=True,
    )
    provider_id = fields.Many2one("phan.he.provider", string="Nhà cung cấp", tracking=True)
    contract_ref = fields.Char(string="Số hợp đồng / mã tham chiếu", tracking=True)
    bandwidth = fields.Char(
        string="Băng thông",
        tracking=True,
        help="Dùng khi khoản chi phí gắn đường truyền / dịch vụ có băng thông.",
    )
    date_start = fields.Date(string="Ngày bắt đầu", tracking=True)
    date_end = fields.Date(string="Ngày kết thúc", tracking=True, index=True)
    cycle_id = fields.Many2one("phan.he.it.renewal.cycle", string="Chu kỳ gia hạn")
    reminder_date = fields.Date(
        string="Ngày nhắc gia hạn",
        compute="_compute_reminder_date",
        store=True,
    )
    amount_untaxed = fields.Monetary(
        string="Cước tháng",
        currency_field="currency_id",
        tracking=True,
        help="Chi phí trước thuế / cước định kỳ.",
    )
    tax_rate = fields.Float(string="Thuế suất (%)", default=10.0)
    amount_tax = fields.Monetary(
        string="Tiền thuế",
        currency_field="currency_id",
        compute="_compute_amounts",
        store=True,
    )
    amount_total = fields.Monetary(
        string="Số tiền thanh toán",
        currency_field="currency_id",
        tracking=True,
        help="Nhập tay — độc lập với cước tháng.",
    )
    amount_untaxed_actual = fields.Monetary(
        string="Chi phí thực tế trước thuế",
        currency_field="currency_id",
        help="Phân biệt với chi phí dự kiến (amount_untaxed).",
    )
    currency_id = fields.Many2one(
        "res.currency",
        string="Đồng tiền",
        default=lambda self: self.env.ref("base.VND", raise_if_not_found=False)
        or self.env.company.currency_id,
    )
    date_due = fields.Date(string="Ngày thanh toán dự kiến", tracking=True)
    date_paid = fields.Date(string="Ngày thanh toán thực tế", tracking=True)
    responsible_id = fields.Many2one("hr.employee", string="Người phụ trách", tracking=True)
    renewal_state = fields.Selection(
        [
            ("active", "Đang hiệu lực"),
            ("expiring", "Sắp hết hạn"),
            ("quote", "Chờ báo giá"),
            ("approval", "Chờ phê duyệt"),
            ("processing", "Đang xử lý gia hạn"),
            ("renewed", "Đã gia hạn"),
            ("overdue", "Quá hạn"),
            ("cancel", "Đã hủy"),
        ],
        string="Trạng thái gia hạn",
        default="active",
        required=True,
        tracking=True,
    )
    payment_state = fields.Selection(
        [
            ("not_due", "Chưa đến hạn thanh toán"),
            ("confirm", "Chờ xác nhận"),
            ("pending", "Chờ thanh toán"),
            ("paid", "Đã thanh toán"),
            ("overdue", "Quá hạn thanh toán"),
            ("cancel", "Đã hủy"),
        ],
        string="Trạng thái thanh toán",
        default="not_due",
        required=True,
        tracking=True,
    )
    note = fields.Text(string="Ghi chú")
    attachment = fields.Binary(string="Hóa đơn", attachment=True)
    attachment_name = fields.Char(string="Tên tệp hóa đơn")
    period_ids = fields.One2many("phan.he.it.cost.period", "cost_id", string="Lịch sử kỳ gia hạn")
    company_id = fields.Many2one(
        "res.company",
        related="asset_id.company_id",
        store=True,
        readonly=True,
    )
    active = fields.Boolean(default=True)
    remaining_days = fields.Integer(compute="_compute_remaining_days", string="Số ngày còn lại")
    remaining_label = fields.Char(
        compute="_compute_remaining_days",
        string="Thời gian còn lại",
    )
    asset_name = fields.Char(
        related="asset_id.name",
        string="Tên thiết bị",
        store=True,
    )

    @api.depends("amount_untaxed", "tax_rate")
    def _compute_amounts(self):
        for rec in self:
            base = float(rec.amount_untaxed or 0.0)
            rec.amount_tax = base * float(rec.tax_rate or 0.0) / 100.0

    @api.depends("date_end", "cycle_id")
    def _compute_reminder_date(self):
        for rec in self:
            if not rec.date_end:
                rec.reminder_date = False
                continue
            # Mặc định nhắc trước 30 ngày (cấu hình chi tiết ở reminder.config)
            rec.reminder_date = rec.date_end - relativedelta(days=30)

    @api.depends("date_end")
    def _compute_remaining_days(self):
        today = fields.Date.context_today(self)
        for rec in self:
            if not rec.date_end:
                rec.remaining_days = 0
                rec.remaining_label = "—"
                continue
            days = (rec.date_end - today).days
            rec.remaining_days = days
            if days > 0:
                rec.remaining_label = "Còn %s ngày" % days
            elif days == 0:
                rec.remaining_label = "Hết hạn hôm nay"
            else:
                rec.remaining_label = "Quá hạn %s ngày" % abs(days)

    def _allowed_renewal_transitions(self):
        return {
            "active": {"expiring", "quote", "overdue", "cancel"},
            "expiring": {"quote", "approval", "processing", "overdue", "cancel"},
            "quote": {"approval", "processing", "cancel"},
            "approval": {"processing", "quote", "cancel"},
            "processing": {"renewed", "cancel"},
            "renewed": {"active", "expiring"},
            "overdue": {"quote", "approval", "processing", "cancel"},
            "cancel": set(),
        }

    def action_set_renewal_state(self, new_state):
        allowed = self._allowed_renewal_transitions()
        for rec in self:
            if new_state not in allowed.get(rec.renewal_state, set()):
                raise UserError(
                    "Không thể chuyển trạng thái gia hạn từ «%s» sang «%s»."
                    % (rec.renewal_state, new_state)
                )
            rec.renewal_state = new_state
        return True

    def action_mark_expiring(self):
        return self.action_set_renewal_state("expiring")

    def action_request_quote(self):
        return self.action_set_renewal_state("quote")

    def action_submit_approval(self):
        return self.action_set_renewal_state("approval")

    def action_start_processing(self):
        return self.action_set_renewal_state("processing")

    def action_mark_payment_confirm(self):
        for rec in self:
            if rec.payment_state in ("cancel", "paid"):
                raise UserError("Không thể chuyển trạng thái thanh toán từ «%s»." % rec.payment_state)
            rec.payment_state = "confirm"
        return True

    def action_mark_paid(self):
        for rec in self:
            if rec.payment_state not in ("confirm", "pending", "overdue"):
                raise UserError("Chỉ xác nhận đã thanh toán khi đang chờ xác nhận/thanh toán.")
            rec.payment_state = "paid"
            if not rec.date_paid:
                rec.date_paid = fields.Date.context_today(self)
        return True

    def action_confirm_renewal(self, date_start=None, date_end=None, amount_untaxed=None, note=None):
        """Gia hạn: lưu kỳ cũ vào period, cập nhật hạn mới — không ghi đè lịch sử."""
        self.ensure_one()
        if self.renewal_state not in ("processing", "approval", "quote", "expiring", "overdue"):
            raise UserError("Chỉ xác nhận gia hạn khi khoản đang trong quy trình xử lý.")
        Period = self.env["phan.he.it.cost.period"]
        Period.create({
            "cost_id": self.id,
            "date_start": self.date_start,
            "date_end": self.date_end,
            "amount_untaxed": self.amount_untaxed,
            "amount_tax": self.amount_tax,
            "amount_total": self.amount_total,
            "payment_state": self.payment_state,
            "date_paid": self.date_paid,
            "note": note or "Kỳ trước khi gia hạn",
        })
        vals = {
            "renewal_state": "renewed",
            "payment_state": "not_due",
            "date_paid": False,
        }
        if date_start:
            vals["date_start"] = date_start
        if date_end:
            vals["date_end"] = date_end
        elif self.cycle_id and self.date_end:
            vals["date_end"] = self.date_end + relativedelta(months=self.cycle_id.months or 12)
            vals["date_start"] = self.date_end + relativedelta(days=1)
        if amount_untaxed is not None:
            vals["amount_untaxed"] = amount_untaxed
        self.write(vals)
        self.renewal_state = "active"
        return True

    _PIE_COLORS = (
        "#2ecc71",
        "#f1c40f",
        "#e74c3c",
        "#95a5a6",
        "#3498db",
        "#9b59b6",
        "#e67e22",
        "#1abc9c",
        "#34495e",
        "#fd79a8",
    )
    _PIE_REPORTS = {
        "month": {
            "title": "Chi phí tháng",
            "slice": "asset",
            "period": "month",
            "dimension": "Phân bổ theo tài sản",
            "filename": "Chi_phi_thang.xlsx",
        },
        "quarter": {
            "title": "Chi phí quý",
            "slice": "month",
            "period": "quarter",
            "dimension": "Phân bổ theo tháng trong quý",
            "filename": "Chi_phi_quy.xlsx",
        },
        "year": {
            "title": "Chi phí năm",
            "slice": "month",
            "period": "year",
            "dimension": "Phân bổ theo tháng",
            "filename": "Chi_phi_nam.xlsx",
        },
        "asset": {
            "title": "Chi phí theo tài sản",
            "slice": "asset",
            "period": None,
            "dimension": "Phân bổ theo tài sản",
            "filename": "Chi_phi_theo_tai_san.xlsx",
        },
        "estimate": {
            "title": "Dự toán chi phí",
            "slice": "asset",
            "period": None,
            "dimension": "Lọc chi phí danh sách thiết bị theo tháng",
            "filename": "Du_toan_chi_phi.xlsx",
        },
        "schedule": {
            "title": "Lịch gia hạn",
            "slice": "asset",
            "period": None,
            "dimension": "Khoản chi hết hạn / cần gia hạn theo tháng",
            "filename": "Lich_gia_han.xlsx",
        },
        "provider": {
            "title": "Chi phí theo nhà cung cấp",
            "slice": "provider",
            "period": None,
            "dimension": "Phân bổ theo nhà cung cấp",
            "filename": "Chi_phi_theo_nha_cung_cap.xlsx",
        },
    }
    _RENEWAL_LABELS = {
        "active": "Đang hiệu lực",
        "expiring": "Sắp hết hạn",
        "quote": "Chờ báo giá",
        "approval": "Chờ phê duyệt",
        "processing": "Đang xử lý gia hạn",
        "renewed": "Đã gia hạn",
        "overdue": "Quá hạn",
        "cancel": "Đã hủy",
    }

    def _pie_report_meta(self, report_key):
        meta = self._PIE_REPORTS.get(report_key)
        if not meta:
            raise UserError("Không có báo cáo này.")
        return meta

    def _pie_period(self, period_key, filters=None):
        today = fields.Date.context_today(self)
        filters = filters or {}
        mode = filters.get("period_mode") or period_key
        year = int(filters.get("year") or today.year)
        if mode == "month":
            month = int(filters.get("month") or today.month)
            month = min(max(month, 1), 12)
            start = today.replace(year=year, month=month, day=1)
            end = start + relativedelta(months=1, days=-1)
            label = "Tháng %s/%s" % (start.month, start.year)
            prev_start = start - relativedelta(months=1)
            prev_end = start - relativedelta(days=1)
            return start, end, label, prev_start, prev_end, "tháng trước"
        if mode == "quarter":
            quarter = int(filters.get("quarter") or ((today.month - 1) // 3 + 1))
            quarter = min(max(quarter, 1), 4)
            start = today.replace(year=year, month=(quarter - 1) * 3 + 1, day=1)
            end = start + relativedelta(months=3, days=-1)
            label = "Quý %s/%s" % (quarter, year)
            prev_start = start - relativedelta(months=3)
            prev_end = start - relativedelta(days=1)
            return start, end, label, prev_start, prev_end, "quý trước"
        if mode == "year":
            start = today.replace(year=year, month=1, day=1)
            end = today.replace(year=year, month=12, day=31)
            label = "Năm %s" % year
            prev_start = start.replace(year=year - 1)
            prev_end = end.replace(year=year - 1)
            return start, end, label, prev_start, prev_end, "năm trước"
        return None, None, "Toàn bộ khoản chi phí", None, None, ""

    def _pie_period_anchor(self, rec):
        """Mốc kỳ cho báo cáo tháng/quý/năm: theo ngày bắt đầu."""
        return rec.date_start or rec.date_end

    def _pie_in_period(self, rec, start, end, mode="anchor"):
        """mode=anchor: date_start nằm trong kỳ.
        mode=overlap: khoản chi còn hiệu lực giao với kỳ (dùng cho theo tài sản / NCC).
        mode=end: ngày hết hạn trong kỳ, hoặc quá hạn trước kỳ (lịch gia hạn).
        """
        if not start:
            return True
        if mode == "overlap":
            ds = rec.date_start or rec.date_end
            de = rec.date_end or rec.date_start
            if not ds:
                return False
            return ds <= end and (not de or de >= start)
        if mode == "end":
            de = rec.date_end
            if not de:
                return False
            # Trong tháng chọn, hoặc quá hạn từ trước (còn theo dõi gia hạn).
            if start <= de <= end:
                return True
            return de < start and (rec.renewal_state or "") in (
                "active",
                "expiring",
                "overdue",
            )
        anchor = self._pie_period_anchor(rec)
        if not anchor:
            return False
        return start <= anchor <= end

    def _pie_slice(self, rec, kind):
        if kind == "category":
            return rec.category_id.id or 0, rec.category_id.name or "Chưa phân hạng mục"
        if kind == "asset":
            return rec.asset_id.id or 0, rec.asset_name or "Tài sản"
        if kind == "provider":
            return rec.provider_id.id or 0, rec.provider_id.name or "Chưa có nhà cung cấp"
        anchor = self._pie_period_anchor(rec)
        if kind == "month":
            if not anchor:
                return 0, "Không rõ tháng"
            return anchor.month, "Tháng %s" % anchor.month
        if not anchor:
            return 0, "Không rõ quý"
        quarter = (anchor.month - 1) // 3 + 1
        return quarter, "Quý %s" % quarter

    def _pie_date(self, value):
        return value.strftime("%d/%m/%Y") if value else ""

    # Báo cáo Chi phí theo tài sản: chỉ Máy chủ vật lý + Thiết bị mạng/bảo mật.
    _ASSET_REPORT_TYPE_CODES = ("physical", "network", "security")

    def _pie_asset_type_ok(self, rec):
        code = (rec.asset_id.asset_type_code or "").strip().lower()
        return code in self._ASSET_REPORT_TYPE_CODES

    def _pie_asset_label(self, rec):
        asset = rec.asset_id
        if asset:
            return (asset.name or asset.code or rec.asset_name or "Tài sản").strip()
        return (rec.asset_name or "Tài sản").strip()

    def _pie_sum_in_period(self, start, end, mode="anchor", asset_report=False):
        if not start:
            return 0.0
        records = self.search([]).filtered(
            lambda rec: self._pie_in_period(rec, start, end, mode)
        )
        if asset_report:
            records = records.filtered(self._pie_asset_type_ok)
        return sum(float(rec.amount_total or 0.0) for rec in records)

    @api.model
    def get_it_cost_pie_report(self, report_key, filters=None):
        """Biểu đồ tròn + phiếu danh sách cho một mục Báo cáo máy chủ."""
        meta = self._pie_report_meta(report_key)
        filters = dict(filters or {})
        asset_like = report_key in ("asset", "estimate", "schedule")
        if report_key in ("asset", "estimate", "schedule", "provider") and not filters.get(
            "period_mode"
        ):
            filters["period_mode"] = "month"
        period_key = meta["period"] or filters.get("period_mode")
        start, end, period_label, prev_start, prev_end, prev_label = self._pie_period(
            period_key, filters
        )
        # Theo tài sản / dự toán / NCC: overlap. Lịch gia hạn: theo ngày hết hạn (+ quá hạn).
        # Tháng / quý / năm báo cáo: theo ngày bắt đầu.
        if report_key == "schedule":
            match_mode = "end"
        elif report_key in ("asset", "estimate", "provider"):
            match_mode = "overlap"
        else:
            match_mode = "anchor"
        records = self.search([]).filtered(
            lambda rec: self._pie_in_period(rec, start, end, match_mode)
        )
        if report_key == "schedule":
            records = records.filtered(
                lambda rec: (rec.renewal_state or "")
                in ("active", "expiring", "overdue", "quote", "approval", "processing")
            )
        if asset_like:
            records = records.filtered(self._pie_asset_type_ok)
        records = records.sorted(
            key=lambda rec: (
                rec.date_end or rec.date_start or fields.Date.from_string("1900-01-01"),
                rec.id,
            )
        )
        buckets = OrderedDict()
        rows = []
        total = 0.0
        asset_ids = set()
        for rec in records:
            amount = float(rec.amount_total or 0.0)
            total += amount
            if asset_like:
                key = rec.asset_id.id or 0
                label = self._pie_asset_label(rec)
            else:
                key, label = self._pie_slice(rec, meta["slice"])
            bucket = buckets.get(key)
            if not bucket:
                bucket = {"key": key, "label": label, "amount": 0.0, "count": 0}
                buckets[key] = bucket
            else:
                # Giữ tên mới nhất từ danh sách tài sản
                bucket["label"] = label
            bucket["amount"] += amount
            bucket["count"] += 1
            if rec.asset_id:
                asset_ids.add(rec.asset_id.id)
            elif label:
                asset_ids.add(label)
            rows.append({
                "id": rec.id,
                "asset_name": self._pie_asset_label(rec) if asset_like else (rec.asset_name or ""),
                "provider": rec.provider_id.name or "",
                "bandwidth": rec.bandwidth or "",
                "date_start": self._pie_date(rec.date_start),
                "date_end": self._pie_date(rec.date_end),
                "amount_untaxed": float(rec.amount_untaxed or 0.0),
                "amount_total": amount,
                "note": rec.note or "",
                "remaining_label": rec.remaining_label or "",
                "renewal_state": rec.renewal_state or "",
                "renewal_label": self._RENEWAL_LABELS.get(rec.renewal_state, rec.renewal_state or ""),
                "slice_label": label,
            })
        slices = sorted(buckets.values(), key=lambda item: item["amount"], reverse=True)
        if len(slices) > 8:
            head = slices[:7]
            rest = slices[7:]
            head.append({
                "key": "other",
                "label": "Khác",
                "amount": sum(item["amount"] for item in rest),
                "count": sum(item["count"] for item in rest),
            })
            slices = head
        palette = (
            ("#2563eb", "#14b8a6", "#f59e0b", "#ef4444", "#0ea5e9", "#22c55e", "#f97316", "#6366f1")
            if asset_like
            else self._PIE_COLORS
        )
        for index, item in enumerate(slices):
            item["color"] = palette[index % len(palette)]
            item["percent"] = round(item["amount"] * 100.0 / total, 1) if total else 0.0
            item["amount"] = round(item["amount"], 0)
        prev_total = (
            self._pie_sum_in_period(
                prev_start,
                prev_end,
                match_mode,
                asset_report=asset_like,
            )
            if prev_start
            else 0.0
        )
        delta_pct = 0.0
        if prev_total:
            delta_pct = round((total - prev_total) / abs(prev_total) * 100.0, 1)
        elif total:
            delta_pct = 100.0
        top = slices[0] if slices else {}
        today = fields.Date.context_today(self)
        if report_key == "estimate":
            subtitle = "Chi phí thiết bị trong %s — chọn tháng/năm để lọc" % period_label
        elif report_key == "schedule":
            subtitle = "Gia hạn trong %s + khoản quá hạn trước đó — chọn tháng/năm để lọc" % (
                period_label
            )
        elif report_key == "asset":
            subtitle = "Theo dõi và phân tích chi phí theo từng tài sản trong hệ thống"
        else:
            subtitle = meta["dimension"]
        return {
            "key": report_key,
            "title": meta["title"],
            "subtitle": subtitle,
            "dimension": meta["dimension"],
            "period": period_label,
            "period_mode": filters.get("period_mode") or meta["period"] or "month",
            "year": int(filters.get("year") or today.year),
            "month": int(filters.get("month") or today.month),
            "quarter": int(filters.get("quarter") or ((today.month - 1) // 3 + 1)),
            "total": round(total, 0),
            "prev_total": round(prev_total, 0),
            "delta_pct": delta_pct,
            "prev_label": prev_label or "",
            "asset_count": len(asset_ids) or len({r["asset_name"] for r in rows if r["asset_name"]}),
            "top_asset_name": top.get("label") or "",
            "top_asset_amount": top.get("amount") or 0,
            "count": len(rows),
            "slices": slices,
            "rows": rows,
        }

    @api.model
    def export_it_cost_pie_excel(self, report_key, filters=None):
        """Kết xuất Excel: tổng hợp biểu đồ tròn và phiếu danh sách."""
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise UserError("Thiếu thư viện openpyxl trên server.") from exc

        data = self.get_it_cost_pie_report(report_key, filters)
        meta = self._pie_report_meta(report_key)
        wb = openpyxl.Workbook()
        thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        header_fill = PatternFill("solid", fgColor="4C1D95")
        header_font = Font(bold=True, color="FFFFFF")
        title_font = Font(bold=True, size=14, color="1E1B4B")
        money = '#,##0'

        summary = wb.active
        summary.title = "Biểu đồ"
        summary["A1"] = data["title"]
        summary["A1"].font = title_font
        summary.merge_cells("A1:D1")
        summary["A2"] = "%s · %s" % (data["period"], data["dimension"])
        summary["A3"] = "Tổng phiếu"
        summary["B3"] = data["count"]
        summary["C3"] = "Tổng tiền"
        summary["D3"] = data["total"]
        summary["D3"].number_format = money
        headers = ["Nhóm", "Số phiếu", "Số tiền thanh toán", "Tỷ lệ (%)"]
        for col, header in enumerate(headers, 1):
            cell = summary.cell(5, col, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin
        for row_index, item in enumerate(data["slices"], 6):
            values = [item["label"], item["count"], item["amount"], item["percent"]]
            for col, value in enumerate(values, 1):
                cell = summary.cell(row_index, col, value)
                cell.border = thin
                if col in (2, 3):
                    cell.number_format = money
                if col == 4:
                    cell.number_format = "0.0"
        for col, width in enumerate((32, 14, 24, 14), 1):
            summary.column_dimensions[get_column_letter(col)].width = width

        detail = wb.create_sheet("Phiếu danh sách")
        detail["A1"] = "Phiếu danh sách — %s" % data["title"]
        detail["A1"].font = title_font
        detail.merge_cells("A1:K1")
        detail_headers = [
            "Tên thiết bị",
            "Nhà cung cấp",
            "Băng thông",
            "Ngày bắt đầu",
            "Ngày kết thúc",
            "Cước tháng",
            "Số tiền thanh toán",
            "Ghi chú",
            "Thời gian còn lại",
            "Trạng thái",
            "Nhóm biểu đồ",
        ]
        for col, header in enumerate(detail_headers, 1):
            cell = detail.cell(3, col, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin
        for row_index, row in enumerate(data["rows"], 4):
            values = [
                row["asset_name"],
                row["provider"],
                row["bandwidth"],
                row["date_start"],
                row["date_end"],
                row["amount_untaxed"],
                row["amount_total"],
                row["note"],
                row["remaining_label"],
                row["renewal_label"],
                row["slice_label"],
            ]
            for col, value in enumerate(values, 1):
                cell = detail.cell(row_index, col, value)
                cell.border = thin
                if col in (6, 7):
                    cell.number_format = money
        for col, width in enumerate((28, 24, 16, 16, 16, 16, 22, 28, 20, 20, 24), 1):
            detail.column_dimensions[get_column_letter(col)].width = width

        buffer = io.BytesIO()
        wb.save(buffer)
        return {
            "filename": meta["filename"],
            "file_base64": base64.b64encode(buffer.getvalue()).decode(),
        }


class PhanHeItCostPeriod(models.Model):
    _name = "phan.he.it.cost.period"
    _description = "Lịch sử kỳ gia hạn chi phí IT"
    _order = "date_end desc, id desc"

    cost_id = fields.Many2one(
        "phan.he.it.cost",
        string="Khoản chi phí",
        required=True,
        ondelete="cascade",
        index=True,
    )
    asset_id = fields.Many2one(related="cost_id.asset_id", store=True, string="Tài sản")
    date_start = fields.Date(string="Từ ngày")
    date_end = fields.Date(string="Đến ngày")
    amount_untaxed = fields.Monetary(string="Trước thuế", currency_field="currency_id")
    amount_tax = fields.Monetary(string="Thuế", currency_field="currency_id")
    amount_total = fields.Monetary(string="Tổng", currency_field="currency_id")
    currency_id = fields.Many2one(related="cost_id.currency_id", store=True)
    payment_state = fields.Char(string="TT thanh toán (kỳ)")
    date_paid = fields.Date(string="Ngày thanh toán")
    note = fields.Char(string="Ghi chú")
