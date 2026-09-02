# -*- coding: utf-8 -*-

import base64
import calendar
import html
import io
import re
from datetime import date, datetime

from odoo import api, fields, models
from odoo.exceptions import AccessError, UserError
from odoo.osv import expression

COUNT_CODES = ("S2", "S4", "S9", "C1", "GS1", "GS2", "GC1", "F8", "F7")
JOB_TITLES = [("NV", "NV"), ("NT", "NT"), ("CHT", "CHT"), ("NVPT", "NVPT"), ("TV", "TV")]
HOUR_MAP = {
    "S2": 6.5,
    "S4": 6.0,
    "S9": 8.0,
    "C1": 7.0,
    "GS1": 10.0,
    "GS2": 8.0,
    "GC1": 10.0,
    "F8": 12.5,
    "F7": 13.0,
    "OFF": 0.0,
    "LE": 0.0,
    "LỄ": 0.0,
}
WD_LABELS = ("T2", "T3", "T4", "T5", "T6", "T7", "CN")


def _job_code(employee):
    raw = (employee.job_title or (employee.job_id.name if employee.job_id else "") or "").upper()
    mapping = (
        ("NVPT", "NVPT"),
        ("CHT", "CHT"),
        ("CỬA HÀNG TRƯỞNG", "CHT"),
        ("CUA HANG TRUONG", "CHT"),
        ("NVPT", "NVPT"),
        ("THỬ VIỆC", "TV"),
        ("THU VIEC", "TV"),
        ("NT", "NT"),
        ("NV", "NV"),
    )
    for needle, code in mapping:
        if needle in raw:
            return code
    return "NV"


def _month_int(value):
    try:
        month = int(value or 0)
    except (TypeError, ValueError):
        month = 0
    return month if month in range(1, 13) else 1


def _clean_shift_code(val):
    if not val:
        return ""
    if isinstance(val, dict):
        return _clean_shift_code(val.get("code") or "")
    if hasattr(val, "code") and val.code:
        return _clean_shift_code(val.code)
    token = re.split(r"[\s\(]", str(val).strip())[0]
    token = token.upper() if token else ""
    if token in WD_LABELS:
        return ""
    return token


def _norm_code(value):
    return _clean_shift_code(value)


def _hours_for(code, catalog=None):
    code = _norm_code(code)
    if not code or code in ("OFF", "NVM", "TV1", "LE", "LỄ"):
        return 0.0
    if code in HOUR_MAP:
        return HOUR_MAP[code]
    if catalog and code in catalog:
        return catalog[code]
    if code.startswith("F"):
        return 12.0
    if code.startswith("GC") or code.startswith("GS"):
        return 10.0
    if code.startswith("S"):
        return 8.0
    if code.startswith("C"):
        return 7.0
    return 0.0


class LinkqMonthlyRoster(models.Model):
    _name = "linkq.monthly.roster"
    _description = "Bảng xếp ca tháng"
    _order = "year desc, month desc, id desc"
    _inherit = ["mail.thread", "lug.menu.access.mixin"]
    _linkq_menu_key = "schedule_main"

    def copy(self, default=None):
        raise UserError("Không được nhân bản bảng xếp ca.")

    def unlink(self):
        if self.env.context.get("skip_roster_unlink_check"):
            return super().unlink()
        if not self.env.su and not (
            self.env.user.has_group("base.group_system")
            or self.env.user.has_group("lug_phan_he.group_linkq_manager")
            or self.env.user.has_group("lug_phan_he.group_phan_he_admin")
        ):
            raise UserError("Bạn không có quyền xóa Bảng xếp ca sau khi đã tạo!")
        return super().unlink()

    @api.model
    def action_hard_delete_permanent(self, schedule_ids, confirm_code):
        if not self.env.user.has_group("base.group_system"):
            raise AccessError("Chỉ có Quản trị viên tối cao (Administrator) mới có quyền xóa vĩnh viễn!")
        code = (confirm_code or "").strip().upper()
        if code not in ("DELETE", "XOA"):
            raise UserError("Mã xác nhận không đúng! Vui lòng nhập 'DELETE' để xác nhận.")
        ids = [int(i) for i in (schedule_ids or []) if i]
        records = self.sudo().browse(ids).exists()
        if not records:
            raise UserError("Vui lòng chọn ít nhất một lịch ca cần xóa vĩnh viễn!")
        count = len(records)
        self.env["linkq.monthly.roster.line"].sudo().search([("roster_id", "in", records.ids)]).unlink()
        records.with_context(active_test=False, skip_roster_unlink_check=True).sudo().unlink()
        return {
            "status": "success",
            "message": "Đã xóa vĩnh viễn %s bảng lịch ca khỏi hệ thống!" % count,
        }

    @api.model
    def get_linkq_sidebar_stats(self):
        rosters = self.search([])
        data = self.env["linkq.store.notification"].get_store_notification_cards("all")
        counts = data.get("counts") or {}
        missing_stores = set()
        for rec in rosters:
            n = rec.days_in_month or 0
            morning = rec.lack_morning or []
            evening = rec.lack_evening or []
            limit = min(n, len(morning), len(evening))
            if rec.store_id and any((morning[i] or evening[i]) for i in range(limit)):
                missing_stores.add(rec.store_id.id)
        return {
            "roster_count": len(rosters),
            "missing_store_count": len(missing_stores),
            "notify_expiring": counts.get("expiring") or 0,
            "notify_locked": counts.get("locked") or 0,
            "notify_reminder": counts.get("reminder") or 0,
            "notify_need": counts.get("need") or 0,
            "notify_info": counts.get("info") or 0,
        }

    @api.model
    def action_save_shift_grid(self, schedule_id, changes):
        roster = self.browse(schedule_id)
        if not roster.exists():
            return {"success": False, "message": "Không tìm thấy bảng ca!"}
        roster.check_access("write")
        Line = self.env["linkq.monthly.roster.line"]
        grouped = {}
        for item in changes or []:
            line_id = item.get("line_id")
            if not line_id:
                continue
            grouped.setdefault(int(line_id), []).append(item)
        for line_id, items in grouped.items():
            line = Line.browse(line_id)
            if not line.exists() or line.roster_id.id != roster.id:
                continue
            plan = dict(line._code_map(line.plan_codes))
            actual = dict(line._code_map(line.actual_codes))
            meta = {}
            for item in items:
                kind = item.get("shift_type") or item.get("kind")
                if kind in ("employee_name", "job_title", "employee_code"):
                    meta[kind] = item.get("value")
                    continue
                day = str(item.get("day") or "")
                if not day.isdigit():
                    continue
                code = _norm_code(item.get("code") or item.get("shift_code") or "")
                if kind == "actual":
                    actual[day] = code
                else:
                    plan[day] = code
            vals = {"plan_codes": plan, "actual_codes": actual}
            if "employee_name" in meta:
                vals["employee_name"] = meta["employee_name"] or ""
            if "job_title" in meta:
                vals["job_title"] = meta["job_title"] or "NV"
            if "employee_code" in meta:
                try:
                    vals["employee_code"] = int(meta["employee_code"] or 0)
                except (TypeError, ValueError):
                    vals["employee_code"] = 0
            line.write(vals)
        return {"success": True}

    name = fields.Char(string="Tên bảng xếp ca", required=True, default="Bảng xếp ca tháng")
    year = fields.Integer(string="Năm", required=True, default=lambda self: fields.Date.context_today(self).year)
    month = fields.Selection(
        [(str(i), f"Tháng {i}") for i in range(1, 13)],
        string="Tháng",
        required=True,
        default=lambda self: str(fields.Date.context_today(self).month),
    )
    store_id = fields.Many2one(
        "hr.store",
        string="Cửa hàng",
        required=True,
        index=True,
        ondelete="restrict",
        tracking=True,
        check_company=True,
        default=lambda self: self._default_store_id(),
    )
    region = fields.Selection(
        related="store_id.mien",
        string="Khu vực / Miền",
        store=True,
        readonly=True,
    )
    total_employees = fields.Integer(string="Nhân sự", compute="_compute_sheet_summary", store=True)
    staff_full_count = fields.Integer(compute="_compute_sheet_summary", store=True)
    staff_part_count = fields.Integer(compute="_compute_sheet_summary", store=True)
    staff_summary_text = fields.Char(compute="_compute_sheet_summary", store=True)
    department_id = fields.Many2one("hr.department", string="Phòng ban", ondelete="set null")
    holiday_map = fields.Char(string="Ghi chú lễ", default=False)
    staff_overview_html = fields.Html(string="Tổng quan công", compute="_compute_staff_overview_html")
    employee_summary_html = fields.Html(
        string="Danh sách nhân sự",
        compute="_compute_employee_summary_html",
        sanitize=False,
    )
    state = fields.Selection(
        [("draft", "Nháp"), ("confirmed", "Đã chốt")],
        default="draft",
        required=True,
        tracking=True,
    )
    line_ids = fields.One2many("linkq.monthly.roster.line", "roster_id", string="Nhân sự")
    line_count = fields.Integer(compute="_compute_line_count")
    days_in_month = fields.Integer(compute="_compute_headers", store=True)
    weekday_json = fields.Json(compute="_compute_headers", store=True)
    holiday_json = fields.Json(compute="_compute_headers", store=True)
    title_display = fields.Char(compute="_compute_name", store=True)
    lack_morning = fields.Json(compute="_compute_lack")
    lack_evening = fields.Json(compute="_compute_lack")
    shift_catalog_json = fields.Json(compute="_compute_shift_catalog")
    company_id = fields.Many2one("res.company", default=lambda self: self.env.company, required=True)

    @api.model
    def _default_store_id(self):
        allowed = self.env.user._linkq_allowed_hr_store_ids()
        if allowed and not self._linkq_bypass():
            return allowed[0]
        return self.env["hr.store"].search(
            [("company_id", "=", self.env.company.id), ("active", "=", True)],
            limit=1,
        ).id

    @api.model
    def _linkq_store_scope_domain(self):
        if self.env.su or self._linkq_bypass():
            return []
        store_ids = self.env.user._linkq_allowed_hr_store_ids()
        if not store_ids:
            return [("id", "=", False)]
        return [("store_id", "in", store_ids)]

    def _linkq_store_forbidden(self):
        domain = self._linkq_store_scope_domain()
        if not domain:
            return self.browse()
        allowed = set(self.env.user._linkq_allowed_hr_store_ids())
        return self.filtered(lambda rec: rec.store_id.id and rec.store_id.id not in allowed)

    def _check_access(self, operation):
        result = super()._check_access(operation)
        if result is not None:
            return result
        if self.env.su or self._linkq_bypass() or not self:
            return None
        forbidden = self._linkq_store_forbidden()
        if forbidden:
            return forbidden, lambda: AccessError(
                "Bạn chỉ được xem bảng xếp ca của cửa hàng mình phụ trách."
            )
        return None

    @api.model
    def _search(self, domain, offset=0, limit=None, order=None, **kwargs):
        extra = self._linkq_store_scope_domain()
        if extra and not kwargs.get("bypass_access"):
            domain = expression.AND([extra, domain or []])
        return super()._search(
            domain, offset=offset, limit=limit, order=order, **kwargs
        )

    @api.model
    def fields_get(self, allfields=None, attributes=None):
        res = super().fields_get(allfields=allfields, attributes=attributes)
        extra = self._linkq_store_scope_domain()
        if extra and "store_id" in res:
            res["store_id"]["domain"] = [
                ("id", "in", self.env.user._linkq_allowed_hr_store_ids())
            ]
        return res

    @api.model_create_multi
    def create(self, vals_list):
        allowed = self.env.user._linkq_allowed_hr_store_ids()
        if allowed and not self.env.su and not self._linkq_bypass():
            for vals in vals_list:
                store_id = vals.get("store_id")
                if store_id and store_id not in allowed:
                    raise AccessError(
                        "Bạn chỉ được tạo bảng xếp ca cho cửa hàng mình phụ trách."
                    )
                if not store_id:
                    vals["store_id"] = allowed[0]
        return super().create(vals_list)

    def write(self, vals):
        allowed = self.env.user._linkq_allowed_hr_store_ids()
        if (
            allowed
            and not self.env.su
            and not self._linkq_bypass()
            and vals.get("store_id")
            and vals["store_id"] not in allowed
        ):
            raise AccessError(
                "Bạn chỉ được sửa bảng xếp ca của cửa hàng mình phụ trách."
            )
        return super().write(vals)

    def _compute_shift_catalog(self):
        codes = self.env["linkq.shift.code"].sudo().search([], order="sequence, code")
        catalog = []
        seen = set()
        for rec in codes:
            code = (rec.code or "").strip().upper()
            if not code or code in seen:
                continue
            seen.add(code)
            catalog.append({
                "code": code,
                "label": code,
                "group": rec.shift_group or "",
                "hours": rec.total_hours or 0.0,
            })
        if "LE" not in seen and "LỄ" not in seen:
            catalog.append({"code": "LE", "label": "LE", "group": "OFF", "hours": 0.0})
        for rec in self:
            rec.shift_catalog_json = catalog

    def init(self):
        cr = self.env.cr
        cr.execute("SELECT to_regclass('hr_store')")
        if not cr.fetchone()[0]:
            return
        cr.execute(
            """
            UPDATE linkq_monthly_roster r
               SET store_id = hs.id
              FROM phan_he_store ps
              JOIN hr_store hs
                ON lower(btrim(coalesce(hs.name->>'en_US', hs.name->>'vi_VN', ''))) = lower(btrim(coalesce(ps.name, '')))
                OR (
                    hs.code IS NOT NULL AND ps.code IS NOT NULL
                    AND lower(btrim(hs.code)) = lower(btrim(ps.code))
                )
             WHERE r.store_id = ps.id
            """
        )

    @api.depends("year", "month", "store_id", "name")
    def _compute_name(self):
        for rec in self:
            rec.title_display = rec.name or rec._default_sheet_name()

    def _default_sheet_name(self):
        self.ensure_one()
        year = self.year or datetime.now().year
        month_n = _month_int(self.month or str(datetime.now().month))
        store = (self.store_id.name or "").strip().upper()
        if store:
            return f"BẢNG XẾP CA {store} THÁNG {month_n} NĂM {year}"
        return f"BẢNG XẾP CA THÁNG {month_n} NĂM {year}"

    @api.onchange("store_id", "month", "year")
    def _onchange_store_load_data(self):
        if not self.year:
            self.year = datetime.now().year
        self.name = self._default_sheet_name()
        if self.store_id and not self.line_ids:
            self.line_ids = [
                (0, 0, self._line_vals_for_employee(emp))
                for emp in self._employees_for_store()
            ]

    @api.depends("line_ids", "line_ids.job_title")
    def _compute_sheet_summary(self):
        part_jobs = {"TV"}
        for rec in self:
            n = len(rec.line_ids)
            part = sum(1 for line in rec.line_ids if (line.job_title or "").upper() in part_jobs)
            full = max(n - part, 0)
            rec.total_employees = n
            rec.staff_full_count = full
            rec.staff_part_count = part
            rec.staff_summary_text = f"{n} Nhân viên ({full} Full, {part} Part)" if n else "0 Nhân viên"

    @api.depends(
        "line_ids",
        "line_ids.employee_id",
        "line_ids.employee_name",
        "line_ids.job_title",
        "line_ids.hour_total",
    )
    def _compute_employee_summary_html(self):
        table_style = (
            "width:100%;border-collapse:collapse;font-size:12px;text-align:center;"
            "border:1px solid #cbd5e1;table-layout:fixed;background:#fff"
        )
        th_title = (
            "background-color:#fff;color:#111;font-weight:800;text-align:center;"
            "padding:4px;border:1px solid #cbd5e1"
        )
        th_col = (
            "background-color:#fff;color:#111;font-weight:600;text-align:center;"
            "padding:4px;border:1px solid #cbd5e1"
        )
        td = "border:1px solid #cbd5e1;padding:4px 6px;background:#fff;color:#111"
        for rec in self:
            head = (
                f"<table style='{table_style}'><thead>"
                f"<tr><th colspan='4' style='{th_title}'>NHÂN SỰ</th></tr>"
                f"<tr>"
                f"<th style='{th_col};width:10%'>STT</th>"
                f"<th style='{th_col};width:45%'>TÊN NV</th>"
                f"<th style='{th_col};width:25%'>CHỨC VỤ</th>"
                f"<th style='{th_col};width:20%'>TỔNG GIỜ</th>"
                f"</tr></thead><tbody>"
            )
            if not rec.line_ids:
                rec.employee_summary_html = (
                    head
                    + f"<tr><td colspan='4' style='{td};text-align:center;color:#64748b'>"
                    "Chưa có nhân sự</td></tr></tbody></table>"
                )
                continue
            rows = []
            for idx, line in enumerate(rec.line_ids, start=1):
                name = html.escape(
                    line.employee_name
                    or (line.employee_id.name if line.employee_id else "")
                    or "N/A"
                )
                role = html.escape(line.job_title or "NV")
                hours = line.hour_total or 0.0
                if abs(hours - round(hours)) < 0.05:
                    hours_txt = str(int(round(hours)))
                else:
                    hours_txt = f"{hours:.1f}".replace(".", ",")
                rows.append(
                    "<tr>"
                    f"<td style='{td};text-align:center'>{idx}</td>"
                    f"<td style='{td};text-align:left'>{name}</td>"
                    f"<td style='{td};text-align:left'>{role}</td>"
                    f"<td style='{td};text-align:right'>{hours_txt}</td>"
                    "</tr>"
                )
            rec.employee_summary_html = head + "".join(rows) + "</tbody></table>"

    def _compute_line_count(self):
        for rec in self:
            rec.line_count = len(rec.line_ids)

    @api.depends(
        "store_id",
        "line_ids.employee_name",
        "line_ids.job_title",
        "line_ids.hour_total",
        "line_ids.shift_summary_text",
    )
    def _compute_staff_overview_html(self):
        for rec in self:
            rows = []
            for line in rec.line_ids:
                name = html.escape(line.employee_name or (line.employee_id.name if line.employee_id else "") or "—")
                job = html.escape(line.job_title or "NV")
                hours = f"{line.hour_total or 0:.1f}"
                detail = html.escape(line.shift_summary_text or "--")
                rows.append(
                    "<tr>"
                    f'<td class="ps-3 fw-bold">{name}</td>'
                    f'<td><span class="badge text-bg-light border">{job}</span></td>'
                    f'<td class="text-center fw-bold" style="color:#6D28D9">{hours} h</td>'
                    f'<td class="text-muted">{detail}</td>'
                    "</tr>"
                )
            body = "".join(rows) or (
                '<tr><td colspan="4" class="text-center text-muted py-3">Chưa có nhân viên</td></tr>'
            )
            rec.staff_overview_html = (
                '<table class="table table-sm table-hover m-0 align-middle o_mm_overview_table">'
                "<thead><tr>"
                "<th class='ps-3'>HỌ VÀ TÊN</th><th>CHỨC VỤ</th>"
                "<th class='text-center'>TỔNG GIỜ</th>"
                "<th>CHI TIẾT CA TRONG THÁNG</th>"
                "</tr></thead>"
                f"<tbody>{body}</tbody></table>"
            )

    def action_add_employee_manual(self):
        return self.action_add_employee()

    @api.depends("year", "month", "holiday_map")
    def _compute_headers(self):
        for rec in self:
            year = rec.year or date.today().year
            month = _month_int(rec.month)
            n_days = calendar.monthrange(year, month)[1]
            rec.days_in_month = n_days
            holidays = {}
            for chunk in (rec.holiday_map or "").split(","):
                if ":" in chunk:
                    day, label = chunk.split(":", 1)
                    day = day.strip()
                    if day.isdigit():
                        holidays[int(day)] = label.strip()
            rec.holiday_json = holidays
            headers = []
            for day in range(1, 32):
                if day <= n_days:
                    wd = date(year, month, day).weekday()
                    label = WD_LABELS[wd]
                    weekend = label in ("T7", "CN")
                    sunday = label == "CN"
                else:
                    label = ""
                    weekend = False
                    sunday = False
                headers.append({
                    "day": day,
                    "wd": label,
                    "weekend": weekend,
                    "sunday": sunday,
                    "holiday": holidays.get(day, ""),
                    "valid": day <= n_days,
                })
            rec.weekday_json = headers

    @api.depends("line_ids.plan_codes", "line_ids.actual_codes", "days_in_month")
    def _compute_lack(self):
        for rec in self:
            n = rec.days_in_month or 31
            morning = []
            evening = []
            for day in range(1, 32):
                if day > n:
                    morning.append(False)
                    evening.append(False)
                    continue
                codes = []
                for line in rec.line_ids:
                    actual = (line.actual_codes or {}).get(str(day)) or (line.actual_codes or {}).get(day)
                    plan = (line.plan_codes or {}).get(str(day)) or (line.plan_codes or {}).get(day)
                    codes.append(_norm_code(actual or plan))
                has_s = any(
                    ((c.startswith("S") and not c.startswith("GS")) or c.startswith("F"))
                    for c in codes
                    if c and c != "OFF"
                )
                has_t = any(
                    c.startswith("C") or c.startswith("GC") or c in ("F7", "F8")
                    for c in codes if c and c != "OFF"
                )
                morning.append(not has_s)
                evening.append(not has_t)
            rec.lack_morning = morning
            rec.lack_evening = evening

    def _employees_for_store(self):
        self.ensure_one()
        Emp = self.env["hr.employee"]
        if not self.store_id:
            return Emp.browse()
        domain = [("active", "=", True)]
        if "store_id" in Emp._fields:
            return Emp.search(domain + [("store_id", "=", self.store_id.id)], order="name")
        Version = self.env["hr.version"]
        if "store_id" in Version._fields:
            versions = Version.search([("store_id", "=", self.store_id.id)])
            emp_ids = versions.mapped("employee_id").ids
            return Emp.search(domain + [("id", "in", emp_ids)], order="name")
        if self.department_id:
            return Emp.search(domain + [("department_id", "=", self.department_id.id)], order="name")
        return Emp.browse()

    def _employee_domain(self):
        employees = self._employees_for_store()
        if employees:
            return [("id", "in", employees.ids)]
        return [("active", "=", True), ("id", "=", 0)]

    def _line_vals_for_employee(self, employee):
        raw = employee.barcode or ""
        emp_code = int(raw) if str(raw).isdigit() else employee.id
        return {
            "employee_id": employee.id,
            "employee_name": employee.name or "",
            "job_title": _job_code(employee),
            "employee_code": emp_code,
            "plan_codes": {},
            "actual_codes": {},
        }

    def action_load_employees(self):
        self.ensure_one()
        employees = self.env["hr.employee"].search(self._employee_domain(), order="name")
        existing = set(self.line_ids.mapped("employee_id").ids)
        vals = []
        for emp in employees:
            if emp.id in existing:
                continue
            vals.append((0, 0, self._line_vals_for_employee(emp)))
        if vals:
            self.write({"line_ids": vals})
        return True

    def action_add_employee(self):
        self.ensure_one()
        self.write({"line_ids": [(0, 0, {
            "job_title": "NV",
            "plan_codes": {},
            "actual_codes": {},
        })]})
        return True

    def action_copy_plan_to_actual(self):
        self.ensure_one()
        for line in self.line_ids:
            line.write({"actual_codes": line._code_map(line.plan_codes)})
        return True

    def action_clean_shift_display(self):
        """Làm sạch mã ca đã lưu kèm giờ (F7 (09:00 - 22:00) → F7)."""
        for sheet in self:
            for line in sheet.line_ids:
                vals = {
                    "plan_codes": line._code_map(line.plan_codes),
                    "actual_codes": line._code_map(line.actual_codes),
                }
                for day in range(1, 32):
                    field_name = f"d_{day}"
                    if field_name in line._fields:
                        raw_val = line[field_name]
                        if raw_val:
                            vals[field_name] = _clean_shift_code(raw_val)
                line.write(vals)
        self.line_ids._compute_counts()
        return True

    def action_recompute_hours(self):
        self.action_clean_shift_display()
        return True

    def action_confirm(self):
        self.write({"state": "confirmed"})

    def action_draft(self):
        self.write({"state": "draft"})

    def action_export_excel(self):
        self.ensure_one()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as err:
            raise UserError("Cần thư viện openpyxl để xuất Excel.") from err

        days = self.days_in_month or calendar.monthrange(self.year or 2026, _month_int(self.month))[1]
        last_col = 4 + days
        headers = self.weekday_json or []
        header_fill = PatternFill("solid", fgColor="EDE9FE")
        title_fill = PatternFill("solid", fgColor="5B21B6")
        info_fill = PatternFill("solid", fgColor="F5F3FF")
        thin_side = Side(style="thin", color="64748B")
        med_side = Side(style="medium", color="4C1D95")
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_mid = Alignment(horizontal="left", vertical="center", wrap_text=True)

        def _box(r1, c1, r2, c2):
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    cell = ws.cell(r, c)
                    cell.border = Border(
                        left=med_side if c == c1 else thin_side,
                        right=med_side if c == c2 else thin_side,
                        top=med_side if r == r1 else thin_side,
                        bottom=med_side if r == r2 else thin_side,
                    )

        wb = Workbook()
        ws = wb.active
        ws.title = f"T{_month_int(self.month):02d}.{str(self.year)[-2:]}"

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        title_cell = ws.cell(1, 1, (self.title_display or self.name or "").upper())
        title_cell.font = Font(bold=True, size=16, color="FFFFFF")
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        for col in range(2, last_col + 1):
            ws.cell(1, col).fill = title_fill

        store_name = self.store_id.display_name or self.store_id.name or ""
        region = self.region or ""
        period = self.period_label or f"Tháng {_month_int(self.month)} / {self.year or ''}"
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        info_cell = ws.cell(
            2,
            1,
            f"Cửa hàng: {store_name}    |    Khu vực: {region}    |    Kỳ xếp ca: {period}    |    Từ ngày 01 đến ngày {days}",
        )
        info_cell.font = Font(size=11, color="4C1D95")
        info_cell.fill = info_fill
        info_cell.alignment = left_mid
        ws.row_dimensions[2].height = 20
        for col in range(2, last_col + 1):
            ws.cell(2, col).fill = info_fill

        staff_bits = []
        for idx, line in enumerate(self.line_ids, 1):
            name = line.employee_name or (line.employee_id.name if line.employee_id else "")
            staff_bits.append(f"{idx}. {name} ({line.job_title or 'NV'}) {line.hour_total or 0}h")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
        staff_cell = ws.cell(3, 1, "Nhân sự: " + ("  •  ".join(staff_bits) if staff_bits else "—"))
        staff_cell.font = Font(size=10, color="334155")
        staff_cell.fill = info_fill
        staff_cell.alignment = left_mid
        ws.row_dimensions[3].height = 18
        for col in range(2, last_col + 1):
            ws.cell(3, col).fill = info_fill

        labels = ["NHÂN VIÊN", "CHỨC VỤ", "MÃ ID", "CA"]
        for col, text in enumerate(labels, 1):
            cell = ws.cell(4, col, text)
            cell.font = Font(bold=True, color="5B21B6")
            cell.fill = header_fill
            cell.alignment = center
        for day in range(1, days + 1):
            cell = ws.cell(4, 4 + day, day)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = center
            info = headers[day - 1] if day - 1 < len(headers) else {}
            wd = ws.cell(5, 4 + day, info.get("wd") or "")
            wd.font = Font(size=9, color="64748B")
            wd.fill = header_fill
            wd.alignment = center
        for col in range(1, 5):
            ws.cell(5, col).fill = header_fill
            ws.cell(5, col).alignment = center

        row = 6
        for line in self.line_ids:
            ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=1)
            ws.merge_cells(start_row=row, start_column=2, end_row=row + 2, end_column=2)
            ws.merge_cells(start_row=row, start_column=3, end_row=row + 2, end_column=3)
            name_cell = ws.cell(row, 1, line.employee_name or (line.employee_id.name if line.employee_id else ""))
            name_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row, 2, line.job_title or "").alignment = center
            ws.cell(row, 3, line.employee_code or "").alignment = center
            ws.cell(row, 4, "DỰ KIẾN").alignment = center
            ws.cell(row + 1, 4, "THỰC TẾ").alignment = center
            ws.cell(row + 2, 4, "SỐ GIỜ").alignment = center
            for day in range(1, days + 1):
                key = str(day)
                hours_val = (line.hour_codes or {}).get(key, "") or None
                ws.cell(row, 4 + day, (line.plan_codes or {}).get(key, "")).alignment = center
                ws.cell(row + 1, 4 + day, (line.actual_codes or {}).get(key, "")).alignment = center
                hour_cell = ws.cell(row + 2, 4 + day, hours_val)
                hour_cell.alignment = center
            row += 3

        last_row = max(row - 1, 5)
        _box(1, 1, last_row, last_col)
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 12
        for col in range(5, last_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = 6

        buf = io.BytesIO()
        wb.save(buf)
        attachment = self.env["ir.attachment"].create({
            "name": f"{self.name or 'bang-xep-ca'}.xlsx",
            "datas": base64.b64encode(buf.getvalue()),
            "res_model": self._name,
            "res_id": self.id,
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=true",
            "target": "new",
        }


class LinkqMonthlyRosterLine(models.Model):
    _name = "linkq.monthly.roster.line"
    _description = "Dòng xếp ca tháng"
    _order = "sequence, id"
    _inherit = ["lug.menu.access.mixin"]
    _linkq_menu_key = "schedule_main"

    roster_id = fields.Many2one("linkq.monthly.roster", required=True, ondelete="cascade", index=True)
    sequence = fields.Integer(default=10)

    @api.model
    def _search(self, domain, offset=0, limit=None, order=None, **kwargs):
        extra = self.env["linkq.monthly.roster"]._linkq_store_scope_domain()
        if extra and not kwargs.get("bypass_access"):
            line_extra = [(f"roster_id.{term[0]}", term[1], term[2]) for term in extra]
            domain = expression.AND([line_extra, domain or []])
        return super()._search(
            domain, offset=offset, limit=limit, order=order, **kwargs
        )
    employee_id = fields.Many2one("hr.employee", string="Nhân viên (HR)", ondelete="set null")
    employee_name = fields.Char(string="Nhân viên")
    job_title = fields.Char(string="Chức vụ", default="NV")
    employee_code = fields.Integer(string="ID")
    plan_codes = fields.Json(string="Dự kiến", default=lambda self: {})
    actual_codes = fields.Json(string="Chính thức", default=lambda self: {})
    plan_notes = fields.Json(string="Ghi chú dự kiến", default=lambda self: {})
    actual_notes = fields.Json(string="Ghi chú thực tế", default=lambda self: {})
    hour_codes = fields.Json(compute="_compute_counts", store=True)
    count_plan = fields.Json(compute="_compute_counts", store=True)
    count_actual = fields.Json(compute="_compute_counts", store=True)
    plan_hours = fields.Float(compute="_compute_counts", store=True)
    actual_hours = fields.Float(compute="_compute_counts", store=True)
    hour_total = fields.Float(compute="_compute_counts", store=True)
    shift_summary_text = fields.Char(compute="_compute_shift_summary_text", store=True)

    _employee_month_uniq = models.Constraint(
        "unique(roster_id, employee_id)",
        "Mỗi nhân viên chỉ có một dòng trong bảng xếp ca tháng.",
    )

    def _notes_map(self, payload):
        data = payload or {}
        result = {}
        for i in range(1, 32):
            item = data.get(str(i)) or data.get(i)
            if not isinstance(item, dict):
                continue
            content = (item.get("content") or "").strip()
            title = (item.get("title") or "").strip()
            if not content and not title:
                continue
            result[str(i)] = {
                "title": title,
                "content": content,
                "author": (item.get("author") or "").strip(),
                "updated_at": item.get("updated_at") or "",
            }
        return result

    def _code_map(self, payload):
        data = payload or {}
        result = {}
        for i in range(1, 32):
            result[str(i)] = _norm_code(data.get(str(i)) or data.get(i) or "")
        return result

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            roster_id = vals.get("roster_id")
            if roster_id and "sequence" not in vals:
                existing = self.search([("roster_id", "=", roster_id)], order="sequence asc, id asc", limit=1)
                vals["sequence"] = (existing.sequence - 1) if existing else 1
            if "plan_codes" in vals:
                vals["plan_codes"] = self._code_map(vals["plan_codes"])
            if "actual_codes" in vals:
                vals["actual_codes"] = self._code_map(vals["actual_codes"])
            if "plan_notes" in vals:
                vals["plan_notes"] = self._notes_map(vals["plan_notes"])
            if "actual_notes" in vals:
                vals["actual_notes"] = self._notes_map(vals["actual_notes"])
        return super().create(vals_list)

    def write(self, vals):
        vals = dict(vals)
        if "plan_codes" in vals:
            vals["plan_codes"] = self._code_map(vals["plan_codes"])
        if "actual_codes" in vals:
            vals["actual_codes"] = self._code_map(vals["actual_codes"])
        if "plan_notes" in vals:
            vals["plan_notes"] = self._notes_map(vals["plan_notes"])
        if "actual_notes" in vals:
            vals["actual_notes"] = self._notes_map(vals["actual_notes"])
        return super().write(vals)

    def read(self, fields=None, load="_classic_read"):
        rows = super().read(fields, load=load)
        for row in rows:
            if "plan_codes" in row:
                row["plan_codes"] = self._code_map(row.get("plan_codes"))
            if "actual_codes" in row:
                row["actual_codes"] = self._code_map(row.get("actual_codes"))
        return rows

    def _count(self, codes):
        tallies = {code: 0 for code in COUNT_CODES}
        for value in codes.values():
            if value in tallies:
                tallies[value] += 1
        return tallies

    def _sum_hours(self, codes, catalog):
        total = 0.0
        for value in codes.values():
            total += _hours_for(value, catalog)
        return total

    def _hour_map(self, plan, actual, catalog):
        hours = {}
        for day in range(1, 32):
            key = str(day)
            code = actual.get(key) or plan.get(key)
            hours[key] = _hours_for(code, catalog) or ""
        return hours

    @api.depends("plan_codes", "actual_codes")
    def _compute_counts(self):
        catalog = {
            rec.code: rec.total_hours
            for rec in self.env["linkq.shift.code"].sudo().search([])
            if rec.code
        }
        for rec in self:
            plan = rec._code_map(rec.plan_codes)
            actual = rec._code_map(rec.actual_codes)
            rec.count_plan = rec._count(plan)
            rec.count_actual = rec._count(actual)
            rec.plan_hours = rec._sum_hours(plan, catalog)
            rec.actual_hours = rec._sum_hours(actual, catalog)
            hours = rec._hour_map(plan, actual, catalog)
            rec.hour_codes = hours
            rec.hour_total = sum(float(v or 0) for v in hours.values())

    @api.depends("count_plan", "count_actual")
    def _compute_shift_summary_text(self):
        for rec in self:
            counts = rec.count_actual or rec.count_plan or {}
            parts = []
            for code in COUNT_CODES:
                n = int(counts.get(code) or 0)
                if n:
                    parts.append(f"{code} ({n}N)")
            rec.shift_summary_text = " • ".join(parts) if parts else "--"
