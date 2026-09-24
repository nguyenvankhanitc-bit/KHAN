# -*- coding: utf-8 -*-

from dateutil.relativedelta import relativedelta

from odoo import api, fields, models


MIEN_COLORS = {
    "BAC": "#2563eb",
    "TRUNG": "#f97316",
    "NAM": "#14b8a6",
    "DTT": "#f97316",
    "VP": "#94a3b8",
}
MIEN_FALLBACK_COLORS = ["#2563eb", "#14b8a6", "#f97316", "#94a3b8", "#7c3aed", "#0d9488"]
TOTAL_SERIES_COLOR = "#7c3aed"


class PhanHeDashboard(models.AbstractModel):
    _name = "phan.he.dashboard"
    _description = "Dashboard quản lý dịch vụ"

    @api.model
    def get_dashboard_data(self, filters=None):
        filters = filters or {}
        today = fields.Date.context_today(self)
        now = fields.Datetime.context_timestamp(self, fields.Datetime.now())

        year = today.year
        if filters.get("date_from"):
            year = fields.Date.to_date(filters["date_from"]).year
        date_from = fields.Date.to_date(filters.get("date_from") or f"{year}-01-01")
        date_to = fields.Date.to_date(filters.get("date_to") or f"{year}-12-31")
        year = date_from.year

        mien_id = int(filters["mien_id"]) if filters.get("mien_id") else False
        area_id = int(filters["area_id"]) if filters.get("area_id") else False
        emp_id = int(filters["employee_id"]) if filters.get("employee_id") else False

        Service = self.env["phan.he.service"]
        Payment = self.env["phan.he.payment"]
        Mien = self.env["phan.he.mien"]

        s_domain = [("active", "=", True)]
        service_type_code = (filters.get("service_type_code") or "").strip().lower()
        if service_type_code:
            stype = self.env["phan.he.service.type"].search(
                [("code", "=", service_type_code)], limit=1
            )
            if stype:
                s_domain.append(("service_type_id", "=", stype.id))
        if mien_id:
            s_domain.append(("mien_id", "=", mien_id))
        if area_id:
            s_domain.append(("area_id", "=", area_id))
        if emp_id:
            s_domain.append(("store_id.responsible_id", "=", emp_id))

        services = Service.search(s_domain)
        miens = Mien.search([("active", "=", True)], order="sequence, name")
        mien_meta = []
        for i, m in enumerate(miens):
            color = MIEN_COLORS.get((m.code or "").upper()) or MIEN_FALLBACK_COLORS[i % len(MIEN_FALLBACK_COLORS)]
            short = (m.name or "").replace("Miền ", "").strip() or m.code or m.name
            mien_meta.append({
                "id": m.id,
                "code": m.code or "",
                "name": m.name or "",
                "short": short,
                "color": color,
            })

        # --- Chi phí tháng theo miền ---
        # Cộng cước tháng của mọi HĐ còn hiệu lực trong tháng (không hủy)
        def services_in_month(recs, m_start, m_end):
            return recs.filtered(
                lambda s: s.state not in ("cancel",)
                and (not s.date_start or s.date_start <= m_end)
                and (not s.date_end or s.date_end >= m_start)
            )

        def cost_by_mien(recs, m_start, m_end):
            """Tổng cước tháng: mọi HĐ hiệu lực trong tháng (không hủy)."""
            active = services_in_month(recs, m_start, m_end)
            result = {m["id"]: 0.0 for m in mien_meta}
            for svc in active:
                mid = svc.mien_id.id
                if mid in result:
                    result[mid] += svc.contract_amount or 0.0
            return result

        cur_month_start = today.replace(day=1)
        cur_month_end = cur_month_start + relativedelta(months=1, days=-1)
        prev_month_start = cur_month_start - relativedelta(months=1)
        prev_month_end = cur_month_start - relativedelta(days=1)

        cur_map = cost_by_mien(services, cur_month_start, cur_month_end)
        prev_map = cost_by_mien(services, prev_month_start, prev_month_end)

        region_kpis = []
        for m in mien_meta:
            cur = cur_map.get(m["id"], 0.0)
            prev = prev_map.get(m["id"], 0.0)
            region_kpis.append({
                **m,
                "amount": cur,
                "delta": self._delta_pct(cur, prev),
                "title": f"CHI PHÍ {m['name'].upper()}",
            })

        # --- Cảnh báo ---
        soon30 = today + relativedelta(days=30)
        expire_soon_svcs = services.filtered(
            lambda s: s.state == "active" and s.date_end and today <= s.date_end <= soon30
        )
        expired_svcs = services.filtered(
            lambda s: s.date_end and s.date_end < today and s.state not in ("cancel", "liquidated")
        )
        due_soon_pays = Payment.search([
            ("payment_state", "in", ("due_soon", "pending", "not_due")),
            ("service_id", "in", services.ids or [0]),
            ("date_due", "!=", False),
            ("date_due", ">=", today),
            ("date_due", "<=", soon30),
        ])

        def breakdown_services(recs):
            counts = {m["id"]: 0 for m in mien_meta}
            for rec in recs:
                mid = rec.mien_id.id
                if mid in counts:
                    counts[mid] += 1
            return [
                {"id": m["id"], "short": m["short"], "count": counts[m["id"]], "color": m["color"]}
                for m in mien_meta
            ]

        def breakdown_payments(pays):
            counts = {m["id"]: 0 for m in mien_meta}
            for pay in pays:
                mid = pay.service_id.mien_id.id
                if mid in counts:
                    counts[mid] += 1
            return [
                {"id": m["id"], "short": m["short"], "count": counts[m["id"]], "color": m["color"]}
                for m in mien_meta
            ]

        alert_cards = [
            {
                "id": "expire_soon",
                "level": "warn",
                "icon": "fa-exclamation-triangle",
                "title": "SẮP HẾT HẠN (30 Ngày)",
                "count": len(expire_soon_svcs),
                "unit": "Hợp đồng",
                "breakdown": breakdown_services(expire_soon_svcs),
                "action": "lug_phan_he.action_phan_he_service_expire_soon",
            },
            {
                "id": "expired",
                "level": "danger",
                "icon": "fa-ban",
                "title": "ĐÃ HẾT HẠN / QUÁ HẠN",
                "count": len(expired_svcs),
                "unit": "Hợp đồng",
                "breakdown": breakdown_services(expired_svcs),
                "action": "lug_phan_he.action_phan_he_service_expired",
            },
            {
                "id": "pay_due_soon",
                "level": "info",
                "icon": "fa-bell",
                "title": "SẮP ĐẾN HẠN THANH TOÁN",
                "count": len(due_soon_pays),
                "unit": "Hợp đồng",
                "breakdown": breakdown_payments(due_soon_pays),
                "action": "lug_phan_he.action_phan_he_payment_due_soon",
            },
        ]

        # --- Biểu đồ 12 tháng (theo năm filter) ---
        trend_months = []
        series = {m["id"]: [] for m in mien_meta}
        totals_row = []
        year_totals = {m["id"]: 0.0 for m in mien_meta}
        grand_total = 0.0

        for month in range(1, 13):
            m_start = fields.Date.to_date(f"{year}-{month:02d}-01")
            m_end = m_start + relativedelta(months=1, days=-1)
            cmap = cost_by_mien(services, m_start, m_end)
            month_total = sum(cmap.values())
            for m in mien_meta:
                val = cmap.get(m["id"], 0.0)
                series[m["id"]].append(val)
                year_totals[m["id"]] += val
            grand_total += month_total

            if month < today.month and year == today.year:
                status = "done"
                status_label = "Hoàn tất"
            elif month == today.month and year == today.year:
                status = "current"
                status_label = "Tháng hiện tại"
            elif year < today.year:
                status = "done"
                status_label = "Hoàn tất"
            elif year > today.year:
                status = "forecast"
                status_label = "Dự kiến"
            else:
                status = "forecast"
                status_label = "Dự kiến"

            totals_row.append({
                "month": month,
                "label": f"Tháng {month:02d}",
                "is_current": status == "current",
                "amounts": [
                    {"mien_id": m["id"], "amount": cmap.get(m["id"], 0.0), "color": m["color"]}
                    for m in mien_meta
                ],
                "total": month_total,
                "status": status,
                "status_label": status_label,
            })
            trend_months.append(f"T{month:02d}")

        max_trend = max(
            [month_total for month_total in (
                sum(series[m["id"]][i] for m in mien_meta) for i in range(12)
            )] + [v for vals in series.values() for v in vals] + [1]
        ) or 1

        # Series "Tổng chi phí" đứng đầu — multi-line chart
        total_values = [
            sum(series[m["id"]][i] for m in mien_meta) for i in range(12)
        ]
        trend_series = [
            {
                "id": "total",
                "name": "Tổng chi phí",
                "short": "Tổng",
                "color": TOTAL_SERIES_COLOR,
                "values": total_values,
                "is_total": True,
            }
        ] + [
            {
                "id": m["id"],
                "name": m["name"],
                "short": m["short"],
                "color": m["color"],
                "values": series[m["id"]],
                "is_total": False,
            }
            for m in mien_meta
        ]

        # Donut / bar: cơ cấu tháng hiện tại
        cur_total = sum(cur_map.values()) or 1.0
        cost_structure = [
            {
                "id": m["id"],
                "name": m["name"],
                "short": m["short"],
                "amount": cur_map.get(m["id"], 0.0),
                "color": m["color"],
                "pct": round((cur_map.get(m["id"], 0.0) / cur_total) * 100, 1),
            }
            for m in mien_meta
            if cur_map.get(m["id"], 0.0) > 0
        ]
        if not cost_structure:
            cost_structure = [
                {
                    "id": m["id"],
                    "name": m["name"],
                    "short": m["short"],
                    "amount": 0.0,
                    "color": m["color"],
                    "pct": 0.0,
                }
                for m in mien_meta
            ]

        # Top 5 chi nhánh / khu vực chi phí cao (tháng hiện tại)
        active_cur = services_in_month(services, cur_month_start, cur_month_end)
        branch_map = {}
        for svc in active_cur:
            area = svc.area_id
            key = area.id if area else (svc.store_id.id or 0)
            name = (area.name if area else (svc.store_id.name or "Khác"))
            row = branch_map.setdefault(key, {"id": key, "name": name, "amount": 0.0})
            row["amount"] += svc.contract_amount or 0.0
        top_branches = sorted(branch_map.values(), key=lambda r: r["amount"], reverse=True)[:5]
        top_max = max((r["amount"] for r in top_branches), default=1) or 1
        for r in top_branches:
            r["pct"] = round((r["amount"] / top_max) * 100, 1)

        # Filter options
        miens_opts = [{"id": m["id"], "name": m["name"]} for m in mien_meta]
        area_domain = [("active", "=", True)]
        if mien_id:
            area_domain.append(("mien_id", "=", mien_id))
        areas = [
            {"id": a.id, "name": a.name, "mien_id": a.mien_id.id}
            for a in self.env["phan.he.area"].search(area_domain, order="name")
        ]
        employees = self.env["hr.employee"].search_read(
            [("id", "in", self.env["phan.he.store"].search([]).mapped("responsible_id").ids)],
            ["id", "name"],
            order="name",
        )

        return {
            "user_name": self.env.user.name,
            "role_label": self._role_label(),
            "updated_at": now.strftime("%H:%M %d/%m/%Y"),
            "app_title": filters.get("app_title") or "Quản lý dịch vụ",
            "service_type_code": service_type_code or "",
            "year": year,
            "current_month": today.month if year == today.year else 0,
            "date_from": fields.Date.to_string(date_from),
            "date_to": fields.Date.to_string(date_to),
            "mien_id": mien_id or "",
            "area_id": area_id or "",
            "employee_id": emp_id or "",
            "miens": miens_opts,
            "areas": areas,
            "employees": employees,
            "mien_meta": mien_meta,
            "region_kpis": region_kpis,
            "alert_cards": alert_cards,
            "alert_count": (
                len(expire_soon_svcs) + len(expired_svcs) + len(due_soon_pays)
            ),
            "expire_soon": len(expire_soon_svcs),
            "overdue_contract": len(expired_svcs),
            "payment_forecast": self.env["phan.he.service"].get_internet_alert_counts().get(
                "payment_forecast", 0
            ),
            "trend_months": trend_months,
            "trend_series": trend_series,
            "max_trend": max_trend,
            "cost_structure": cost_structure,
            "top_branches": top_branches,
            "monthly_table": totals_row,
            "year_totals": [
                {"mien_id": m["id"], "amount": year_totals[m["id"]], "color": m["color"]}
                for m in mien_meta
            ],
            "year_grand_total": grand_total,
            "currency_symbol": (
                self.env.ref("base.VND", raise_if_not_found=False)
                or self.env["res.currency"].search([("name", "=", "VND")], limit=1)
                or self.env.company.currency_id
            ).symbol or "đ",
        }

    @api.model
    def _delta_pct(self, current, previous):
        if not previous:
            return 100.0 if current else 0.0
        return round(((current - previous) / abs(previous)) * 100, 1)

    @api.model
    def export_monthly_cost_excel(self, filters=None):
        """Xuất bảng tổng hợp chi phí 12 tháng ra Excel."""
        import base64
        import io

        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError as exc:
            from odoo.exceptions import UserError
            raise UserError("Thiếu thư viện openpyxl trên server.") from exc

        data = self.get_dashboard_data(filters or {})
        year = data.get("year") or ""
        mien_meta = data.get("mien_meta") or []
        monthly = data.get("monthly_table") or []
        year_totals = data.get("year_totals") or []

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Chi phi {year}"

        header_fill = PatternFill("solid", fgColor="1E3A5F")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        total_fill = PatternFill("solid", fgColor="E2E8F0")
        current_fill = PatternFill("solid", fgColor="FFF7ED")
        thin = Border(
            left=Side(style="thin", color="94A3B8"),
            right=Side(style="thin", color="94A3B8"),
            top=Side(style="thin", color="94A3B8"),
            bottom=Side(style="thin", color="94A3B8"),
        )
        right = Alignment(horizontal="right", vertical="center")
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        title = f"Bảng tổng hợp chi phí chi tiết từ tháng 1 đến tháng 12 (VNĐ · Năm {year})"
        headers = ["Tháng"] + [f"Chi phí {m.get('name') or ''}" for m in mien_meta] + [
            "Tổng chi phí",
            "Trạng thái / Lưu ý",
        ]
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        cell = ws.cell(1, 1, title)
        cell.font = Font(bold=True, size=13, color="0F172A")
        cell.alignment = left

        for col, text in enumerate(headers, 1):
            c = ws.cell(3, col, text)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center if col == 1 or col == len(headers) else right
            c.border = thin

        row_idx = 4
        for row in monthly:
            values = [row.get("label") or ""]
            for amt in row.get("amounts") or []:
                values.append(float(amt.get("amount") or 0))
            values.append(float(row.get("total") or 0))
            values.append(row.get("status_label") or "")
            for col, val in enumerate(values, 1):
                c = ws.cell(row_idx, col, val)
                c.border = thin
                if col == 1:
                    c.alignment = left
                elif col == len(values):
                    c.alignment = center
                else:
                    c.alignment = right
                    c.number_format = "#,##0"
                if row.get("is_current"):
                    c.fill = current_fill
            row_idx += 1

        # Lũy kế
        foot = ["LŨY KẾ"]
        yt_map = {yt.get("mien_id"): float(yt.get("amount") or 0) for yt in year_totals}
        for m in mien_meta:
            foot.append(yt_map.get(m.get("id"), 0.0))
        foot.append(float(data.get("year_grand_total") or 0))
        foot.append("Tổng cả năm")
        for col, val in enumerate(foot, 1):
            c = ws.cell(row_idx, col, val)
            c.border = thin
            c.fill = total_fill
            c.font = Font(bold=True)
            if col == 1:
                c.alignment = left
            elif col == len(foot):
                c.alignment = center
            else:
                c.alignment = right
                c.number_format = "#,##0"

        ws.column_dimensions["A"].width = 14
        for i in range(2, len(headers)):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
        ws.column_dimensions[openpyxl.utils.get_column_letter(len(headers))].width = 18

        buf = io.BytesIO()
        wb.save(buf)
        return {
            "file_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
            "filename": f"Bang_tong_hop_chi_phi_{year}.xlsx",
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }

    @api.model
    def get_payment_status_report(self, filters=None):
        """Báo cáo 4 danh sách thanh toán theo trạng thái."""
        filters = filters or {}
        today = fields.Date.context_today(self)

        year = today.year
        if filters.get("date_from"):
            year = fields.Date.to_date(filters["date_from"]).year
        date_from = fields.Date.to_date(filters.get("date_from") or f"{year}-01-01")
        date_to = fields.Date.to_date(filters.get("date_to") or f"{year}-12-31")

        mien_id = int(filters["mien_id"]) if filters.get("mien_id") else False
        area_id = int(filters["area_id"]) if filters.get("area_id") else False
        emp_id = int(filters["employee_id"]) if filters.get("employee_id") else False

        Service = self.env["phan.he.service"]
        Payment = self.env["phan.he.payment"]

        s_domain = [("active", "=", True)]
        service_type_code = (filters.get("service_type_code") or "").strip().lower()
        if service_type_code:
            stype = self.env["phan.he.service.type"].search(
                [("code", "=", service_type_code)], limit=1
            )
            if stype:
                s_domain.append(("service_type_id", "=", stype.id))
        if mien_id:
            s_domain.append(("mien_id", "=", mien_id))
        if area_id:
            s_domain.append(("area_id", "=", area_id))
        if emp_id:
            s_domain.append(("store_id.responsible_id", "=", emp_id))

        services = Service.search(s_domain)
        service_ids = services.ids or [0]

        def serialize(pays):
            rows = []
            for p in pays:
                rows.append({
                    "id": p.id,
                    "code": p.code or "",
                    "store": p.store_name or (p.store_id.name if p.store_id else "") or "",
                    "mien": p.mien_id.name if p.mien_id else "",
                    "provider": p.provider_id.name if p.provider_id else "",
                    "period": p.period or "",
                    "invoice_number": p.invoice_number or "",
                    "amount": p.amount or 0.0,
                    "date_due": fields.Date.to_string(p.date_due) if p.date_due else "",
                    "date_paid": fields.Date.to_string(p.date_paid) if p.date_paid else "",
                    "payment_state": p.payment_state,
                })
            return rows

        def search_state(states, order="date_due desc, id desc", limit=200):
            domain = [
                ("service_id", "in", service_ids),
                ("payment_state", "in", list(states)),
            ]
            # Lọc theo kỳ hạn / ngày TT trong khoảng năm đang xem
            domain += [
                "|",
                "&", ("date_due", ">=", date_from), ("date_due", "<=", date_to),
                "&", ("date_paid", ">=", date_from), ("date_paid", "<=", date_to),
            ]
            return Payment.search(domain, order=order, limit=limit)

        paid = search_state(("paid",), order="date_paid desc, id desc")
        due_soon = search_state(("due_soon",), order="date_due asc, id asc")
        pending = search_state(("pending", "not_due", "draft"), order="date_due asc, id asc")
        overdue = search_state(("overdue",), order="date_due asc, id asc")

        def bucket(key, title, tone, pays):
            total = sum(pays.mapped("amount"))
            return {
                "key": key,
                "title": title,
                "tone": tone,
                "count": len(pays),
                "total": total,
                "rows": serialize(pays),
            }

        tables = [
            bucket("paid", "Bảng 1: Danh sách đã thanh toán", "paid", paid),
            bucket("due_soon", "Bảng 2: Danh sách sắp thanh toán", "due_soon", due_soon),
            bucket("pending", "Bảng 3: Danh sách chờ thanh toán", "pending", pending),
            bucket("overdue", "Bảng 4: Danh sách quá hạn", "overdue", overdue),
        ]
        return {
            "date_from": fields.Date.to_string(date_from),
            "date_to": fields.Date.to_string(date_to),
            "tables": tables,
            "grand_count": sum(t["count"] for t in tables),
            "grand_total": sum(t["total"] for t in tables),
        }

    @api.model
    def _role_label(self):
        user = self.env.user
        if user.has_group("lug_phan_he.group_phan_he_admin"):
            return "Administrator"
        if user.has_group("lug_phan_he.group_phan_he_service_manager"):
            return "Service Manager"
        if user.has_group("lug_phan_he.group_phan_he_regional"):
            return "Regional Manager"
        if user.has_group("lug_phan_he.group_phan_he_area"):
            return "Area Manager"
        if user.has_group("lug_phan_he.group_phan_he_staff"):
            return "Staff"
        return "Viewer"

    @api.model
    def get_month_cost_board(self, params=None):
        """Chi phí tháng N — cùng tập Danh sách thanh toán:

        - HĐ cần TT: date_end còn ≤ 30 ngày hoặc quá hạn (need domain)
        - Tổng chi phí = next_payment_amount / cước tháng
        - Đã thanh toán = xác nhận paid gắn kỳ
        - Chưa thanh toán = Tổng − Đã thanh toán
        """
        params = params or {}
        today = fields.Date.context_today(self)
        year = int(params.get("year") or today.year)
        month = int(params.get("month") or today.month)
        month = min(12, max(1, month))
        region = (params.get("region") or "all") or "all"
        provider_id = int(params["provider_id"]) if params.get("provider_id") else False
        pay_filter = (params.get("pay_status") or "all") or "all"

        Service = self.env["phan.he.service"]
        start, end, as_of, soon = Service._month_ky_bounds(year, month)

        extra = []
        if region and region != "all":
            extra.append(("store_mien", "=", region))
        if provider_id:
            extra.append(("provider_id", "=", provider_id))

        services = Service.search_month_cost_ky(year, month, extra_domain=extra or None)

        def period_amount(rec):
            return Service._month_ky_period_amount(rec)

        def paid_amount_for(rec):
            return Service._month_ky_paid_amount(rec, start, end, as_of, soon)

        rows = []
        due_buckets = {"lt7": 0, "d7_15": 0, "d16_30": 0, "overdue": 0}
        status_counts = {"paid": 0, "partial": 0, "unpaid": 0, "none": 0}
        kpi = {
            "total_amount": 0.0,
            "total_count": 0,
            "paid_amount": 0.0,
            "paid_count": 0,
            "unpaid_amount": 0.0,
            "unpaid_count": 0,
        }

        def week_ranges():
            last = end.day
            cuts = [(1, min(5, last)), (6, min(12, last)), (13, min(19, last)), (20, last)]
            out = []
            for i, (a, b) in enumerate(cuts, 1):
                if a > last:
                    break
                b = max(a, b)
                out.append({
                    "key": f"w{i}",
                    "label": f"Tuần {i}",
                    "sub": f"{a:02d}-{b:02d}/{month:02d}",
                    "d0": start.replace(day=a),
                    "d1": start.replace(day=b),
                })
            return out

        weeks = week_ranges()
        week_total = [0.0] * len(weeks)
        week_paid = [0.0] * len(weeks)
        week_remain = [0.0] * len(weeks)

        def week_index(day):
            if not day:
                return None
            for i, w in enumerate(weeks):
                if w["d0"] <= day <= w["d1"]:
                    return i
            return None

        for rec in services:
            monthly = period_amount(rec)
            paid_sum = paid_amount_for(rec)
            # Gắn paid vào kỳ: không vượt số tiền kỳ
            paid_cap = min(paid_sum, monthly) if monthly > 0 else paid_sum
            remain_amt = max(monthly - paid_cap, 0.0)

            if monthly <= 0 and paid_cap <= 0:
                pay_status = "none"
            elif remain_amt <= 0 and (monthly > 0 or paid_cap > 0):
                pay_status = "paid"
            elif paid_cap > 0:
                pay_status = "partial"
            else:
                pay_status = "unpaid"

            if pay_filter not in ("all", "", False) and pay_status != pay_filter:
                continue

            days = (rec.date_end - as_of).days if rec.date_end else 0
            due_key = ""
            if days < 0:
                due_key = "overdue"
                due_buckets["overdue"] += 1
                # Gom quá hạn vào tab "< 7 ngày" để vẫn lọc được trên UI cũ
                due_buckets["lt7"] += 1
            elif 0 <= days < 7:
                due_key = "lt7"
                due_buckets["lt7"] += 1
            elif 7 <= days <= 15:
                due_key = "d7_15"
                due_buckets["d7_15"] += 1
            elif 16 <= days <= 30:
                due_key = "d16_30"
                due_buckets["d16_30"] += 1

            status_counts[pay_status] = status_counts.get(pay_status, 0) + 1
            kpi["total_amount"] += monthly
            kpi["total_count"] += 1
            kpi["paid_amount"] += paid_cap
            if pay_status == "paid":
                kpi["paid_count"] += 1
            elif pay_status in ("unpaid", "partial"):
                kpi["unpaid_count"] += 1

            wi = week_index(rec.date_end) if rec.date_end and start <= rec.date_end <= end else None
            if wi is None and rec.date_end and rec.date_end < start:
                wi = 0  # quá hạn trước tháng → tuần 1
            if wi is None:
                wi = week_index(rec.next_payment_date) if rec.next_payment_date else len(weeks) - 1
            if wi is None:
                wi = len(weeks) - 1
            week_total[wi] += monthly
            week_paid[wi] += paid_cap
            week_remain[wi] += remain_amt

            store = rec.store_id
            provider = rec.provider_id
            rows.append({
                "id": rec.id,
                "store": store.name or rec.name or "—",
                "code": rec.customer_code or rec.code or "",
                "provider": provider.name or "—",
                "provider_id": provider.id or False,
                "region": rec.store_mien or "",
                "date_start": fields.Date.to_string(rec.date_start) if rec.date_start else False,
                "date_end": fields.Date.to_string(rec.date_end) if rec.date_end else False,
                "remaining_days": days,
                "due_key": "lt7" if due_key == "overdue" else due_key,
                "monthly": monthly,
                "paid": paid_cap,
                "remain": remain_amt,
                "pay_status": pay_status,
            })

        # Chưa thanh toán = Tổng − Đã thanh toán (công thức chốt)
        kpi["unpaid_amount"] = max(kpi["total_amount"] - kpi["paid_amount"], 0.0)
        kpi["paid_count"] = status_counts["paid"]
        kpi["unpaid_count"] = status_counts["unpaid"] + status_counts["partial"]

        total_n = kpi["total_count"] or 1
        donut = [
            {"id": "paid", "label": "Đã thanh toán", "count": status_counts["paid"],
             "pct": round(status_counts["paid"] * 100 / total_n), "color": "#22c55e"},
            {"id": "partial", "label": "Thanh toán 1 phần", "count": status_counts["partial"],
             "pct": round(status_counts["partial"] * 100 / total_n), "color": "#f59e0b"},
            {"id": "unpaid", "label": "Chưa thanh toán", "count": status_counts["unpaid"],
             "pct": round(status_counts["unpaid"] * 100 / total_n), "color": "#ef4444"},
        ]

        providers = self.env["phan.he.provider"].search_read(
            [("active", "=", True)], ["name"], order="name asc"
        )
        regions = sorted({r["region"] for r in rows if r.get("region")})

        return {
            "year": year,
            "month": month,
            "month_label": f"Tháng {month:02d}/{year}",
            "as_of": fields.Date.to_string(as_of),
            "kpi": kpi,
            "weeks": [
                {
                    "label": w["label"],
                    "sub": w["sub"],
                    "total": week_total[i],
                    "paid": week_paid[i],
                    "remain": week_remain[i],
                }
                for i, w in enumerate(weeks)
            ],
            "due_bars": [
                {"id": "lt7", "label": "< 7 ngày", "count": due_buckets["lt7"], "color": "#ef4444"},
                {"id": "d7_15", "label": "7 - 15 ngày", "count": due_buckets["d7_15"], "color": "#f97316"},
                {"id": "d16_30", "label": "16 - 30 ngày", "count": due_buckets["d16_30"], "color": "#facc15"},
            ],
            "donut": donut,
            "rows": rows,
            "providers": providers,
            "regions": regions,
        }

    @api.model
    def get_quarter_cost_board(self, params=None):
        """Dashboard Chi phí quý — KPI 4 quý, biểu đồ tháng/NCC/khu vực, danh sách hợp đồng."""
        params = params or {}
        today = fields.Date.context_today(self)
        year = int(params.get("year") or today.year)
        quarter = int(params.get("quarter") or ((today.month - 1) // 3 + 1))
        quarter = min(4, max(1, quarter))
        region = (params.get("region") or "all") or "all"
        provider_id = int(params["provider_id"]) if params.get("provider_id") else False
        pay_filter = (params.get("pay_status") or "all") or "all"

        Service = self.env["phan.he.service"]
        month_names = {1: "01", 2: "02", 3: "03", 4: "04", 5: "05", 6: "06",
                       7: "07", 8: "08", 9: "09", 10: "10", 11: "11", 12: "12"}
        region_label = {
            "Bắc": "Miền Bắc",
            "Nam": "Miền Nam",
            "ĐTT": "Miền Trung",
            "Trung": "Miền Trung",
        }
        region_color = {
            "Miền Bắc": "#3b82f6",
            "Miền Trung": "#22c55e",
            "Miền Nam": "#f59e0b",
            "Tây Nguyên": "#ef4444",
            "Khác": "#94a3b8",
        }

        def quarter_bounds(y, q):
            start_m = (q - 1) * 3 + 1
            start = fields.Date.from_string(f"{y:04d}-{start_m:02d}-01")
            end = (start + relativedelta(months=3)) - relativedelta(days=1)
            return start, end

        def overlap_domain(start, end):
            domain = [
                ("active", "=", True),
                ("service_type_id.code", "=", "internet"),
                ("state", "not in", ("cancel", "draft")),
                "|", ("date_start", "=", False), ("date_start", "<=", end),
                "|", ("date_end", "=", False), ("date_end", ">=", start),
            ]
            if region and region != "all":
                domain.append(("store_mien", "=", region))
            if provider_id:
                domain.append(("provider_id", "=", provider_id))
            return domain

        def pay_of(rec):
            monthly = float(rec.contract_amount or 0.0)
            paid_sum = sum(
                float(p.amount or 0.0)
                for p in rec.payment_ids
                if p.payment_state == "paid"
            )
            remain_amt = max(monthly - paid_sum, 0.0)
            if monthly <= 0 and paid_sum <= 0:
                status = "none"
            elif remain_amt <= 0 and (monthly > 0 or paid_sum > 0):
                status = "paid"
            elif paid_sum > 0:
                status = "partial"
            else:
                status = "unpaid"
            return monthly, paid_sum, remain_amt, status

        def provider_color(name):
            u = (name or "").upper()
            if "VIETTEL" in u:
                return "#4f7cff"
            if "FPT" in u:
                return "#3b82f6"
            if "VNPT" in u:
                return "#22c55e"
            return "#f59e0b"

        quarter_kpis = []
        for q in (1, 2, 3, 4):
            q_start, q_end = quarter_bounds(year, q)
            recs = Service.search(overlap_domain(q_start, q_end))
            providers = set()
            total = 0.0
            count = 0
            for rec in recs:
                monthly, paid_sum, remain_amt, status = pay_of(rec)
                if pay_filter not in ("all", "", False) and status != pay_filter:
                    continue
                total += monthly
                count += 1
                if rec.provider_id:
                    providers.add(rec.provider_id.id)
            sm = (q - 1) * 3 + 1
            quarter_kpis.append({
                "quarter": q,
                "year": year,
                "month_range": f"Tháng {month_names[sm]} – {month_names[sm + 2]}/{year}",
                "total_amount": total,
                "contract_count": count,
                "provider_count": len(providers),
            })

        q_start, q_end = quarter_bounds(year, quarter)
        services = Service.search(overlap_domain(q_start, q_end), order="date_end asc, id asc")
        rows = []
        monthly_map = {}
        provider_map = {}
        region_map = {}
        kpi = {
            "total_amount": 0.0,
            "paid_amount": 0.0,
            "unpaid_amount": 0.0,
            "total_count": 0,
        }
        for rec in services:
            monthly, paid_sum, remain_amt, status = pay_of(rec)
            if pay_filter not in ("all", "", False) and status != pay_filter:
                continue
            kpi["total_amount"] += monthly
            kpi["paid_amount"] += min(paid_sum, monthly) if monthly else paid_sum
            kpi["unpaid_amount"] += remain_amt
            kpi["total_count"] += 1

            bucket_month = rec.date_start.month if rec.date_start and q_start <= rec.date_start <= q_end else (
                rec.date_end.month if rec.date_end and q_start <= rec.date_end <= q_end else q_start.month
            )
            monthly_map[bucket_month] = monthly_map.get(bucket_month, 0.0) + monthly
            pname = rec.provider_id.name or "Khác"
            provider_map[pname] = provider_map.get(pname, 0.0) + monthly
            rkey = region_label.get(rec.store_mien or "", "Khác")
            region_map[rkey] = region_map.get(rkey, 0.0) + monthly

            store = rec.store_id
            provider = rec.provider_id
            rows.append({
                "id": rec.id,
                "store": store.name or rec.name or "—",
                "code": rec.customer_code or rec.code or "",
                "customer_code": rec.customer_code or "",
                "contract_code": rec.code or "",
                "provider": provider.name or "—",
                "provider_id": provider.id or False,
                "bandwidth": rec.bandwidth or "—",
                "region": rec.store_mien or "",
                "date_start": fields.Date.to_string(rec.date_start) if rec.date_start else False,
                "date_end": fields.Date.to_string(rec.date_end) if rec.date_end else False,
                "monthly": monthly,
                "paid": float(rec.next_payment_amount or 0.0),
                "next_payment_amount": float(rec.next_payment_amount or 0.0),
                "remain": remain_amt,
                "pay_status": status,
                "ops_status": rec.ops_status or rec.state or "active",
                "remaining_days": rec.remaining_days if rec.remaining_days is not False else 0,
                "remaining_time": rec.remaining_time or "",
                "alert_level": rec.alert_level or "ok",
            })

        months_in_q = [(quarter - 1) * 3 + i for i in (1, 2, 3)]
        monthly_trend = [
            {
                "label": f"Tháng {m}/{year}",
                "value": monthly_map.get(m, 0.0),
            }
            for m in months_in_q
        ]
        by_provider = [
            {"label": name, "value": amt, "color": provider_color(name)}
            for name, amt in sorted(provider_map.items(), key=lambda x: -x[1])
        ]
        region_total = sum(region_map.values()) or 1.0
        by_region = [
            {
                "label": name,
                "value": amt,
                "percent": round(amt * 100 / region_total),
                "color": region_color.get(name, "#94a3b8"),
            }
            for name, amt in sorted(region_map.items(), key=lambda x: -x[1])
        ]
        providers = self.env["phan.he.provider"].search_read(
            [("active", "=", True)], ["name"], order="name asc"
        )
        regions = sorted({r["region"] for r in rows if r.get("region")})
        sm = (quarter - 1) * 3 + 1
        return {
            "year": year,
            "quarter": quarter,
            "quarter_label": f"Quý {quarter}/{year} (Tháng {month_names[sm]} - {month_names[sm + 2]}/{year})",
            "kpi": kpi,
            "quarter_kpis": quarter_kpis,
            "monthly_trend": monthly_trend,
            "by_provider": by_provider,
            "by_region": by_region,
            "rows": rows,
            "providers": providers,
            "regions": regions,
        }

    @api.model
    def export_quarter_cost_excel(self, params=None):
        """Xuất danh sách chi phí quý ra Excel."""
        import base64
        import io

        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            from odoo.exceptions import UserError

            raise UserError("Thiếu thư viện openpyxl trên server.") from exc

        params = params or {}
        data = self.get_quarter_cost_board(params)
        search = (params.get("search") or "").strip().lower()
        table_region = params.get("table_region") or "all"
        table_provider = str(params.get("table_provider") or "")
        table_status = params.get("table_status") or "all"

        def keep(row):
            if table_region not in ("all", "", None) and row.get("region") != table_region:
                return False
            if table_provider and str(row.get("provider_id") or "") != table_provider:
                return False
            days = int(row.get("remaining_days") or 0)
            alert = row.get("alert_level") or ""
            if table_status == "active" and (days < 0 or alert in ("expired", "danger")):
                return False
            if table_status == "expire_soon" and not (0 <= days <= 30 or alert == "warn"):
                return False
            if table_status == "expired" and not (days < 0 or alert in ("expired", "danger")):
                return False
            if search:
                blob = " ".join([
                    str(row.get("store") or ""),
                    str(row.get("code") or ""),
                    str(row.get("customer_code") or ""),
                    str(row.get("contract_code") or ""),
                    str(row.get("region") or ""),
                    str(row.get("provider") or ""),
                ]).lower()
                if search not in blob:
                    return False
            return True

        rows = [r for r in (data.get("rows") or []) if keep(r)]
        year = data.get("year")
        quarter = data.get("quarter")
        ops_label = {
            "active": "Đang hoạt động",
            "suspend": "Tạm ngưng",
            "liquidated": "Thanh lý",
        }

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Chi phi Q{quarter}-{year}"[:31]

        headers = [
            "STT", "Cửa hàng", "Mã KH", "Mã hợp đồng", "Nhà cung cấp", "Băng thông",
            "Khu vực", "Ngày bắt đầu", "Ngày kết thúc", "Cước tháng",
            "Số tiền thanh toán", "Trạng thái", "Thời gian còn lại",
        ]
        header_fill = PatternFill("solid", fgColor="6D28D9")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        total_fill = PatternFill("solid", fgColor="EDE9FE")
        thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        money_fmt = '#,##0" đ"'
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        tcell = ws.cell(1, 1, f"Danh sách chi phí quý {quarter}/{year}")
        tcell.font = Font(bold=True, size=14, color="4C1D95")
        tcell.alignment = Alignment(horizontal="left", vertical="center")

        for col, text in enumerate(headers, 1):
            c = ws.cell(3, col, text)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin

        def fmt_date(val):
            if not val:
                return ""
            try:
                y, m, d = str(val)[:10].split("-")
                return f"{d}/{m}/{y}"
            except Exception:
                return str(val)

        tot_m = tot_p = 0.0
        for idx, row in enumerate(rows, 1):
            days = int(row.get("remaining_days") or 0)
            remain_txt = row.get("remaining_time") or (
                f"Quá hạn {abs(days)} ngày" if days < 0 else f"Còn {days} ngày"
            )
            ops = row.get("ops_status") or "active"
            status_txt = ops_label.get(ops, ops)
            values = [
                idx,
                row.get("store") or "",
                row.get("customer_code") or row.get("code") or "",
                row.get("contract_code") or "",
                row.get("provider") or "",
                row.get("bandwidth") or "",
                row.get("region") or "",
                fmt_date(row.get("date_start")),
                fmt_date(row.get("date_end")),
                float(row.get("monthly") or 0),
                float(row.get("next_payment_amount") or row.get("paid") or 0),
                status_txt,
                remain_txt,
            ]
            tot_m += values[9]
            tot_p += values[10]
            ridx = 3 + idx
            for col, val in enumerate(values, 1):
                c = ws.cell(ridx, col, val)
                c.border = thin
                c.alignment = Alignment(
                    vertical="center",
                    horizontal="right" if col in (10, 11) else "left",
                )
                if col in (10, 11):
                    c.number_format = money_fmt

        foot_row = 4 + len(rows)
        foot = ["", "TỔNG CỘNG", "", "", "", "", "", "", "", tot_m, tot_p, "", ""]
        for col, val in enumerate(foot, 1):
            c = ws.cell(foot_row, col, val)
            c.border = thin
            c.fill = total_fill
            c.font = Font(bold=True)
            if col in (10, 11):
                c.number_format = money_fmt
                c.alignment = Alignment(horizontal="right")

        for col in range(1, len(headers) + 1):
            maxlen = len(str(headers[col - 1]))
            for r in range(3, foot_row + 1):
                maxlen = max(maxlen, len(str(ws.cell(r, col).value or "")))
            ws.column_dimensions[get_column_letter(col)].width = min(42, max(12, maxlen + 2))

        buf = io.BytesIO()
        wb.save(buf)
        return {
            "file_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
            "filename": f"Chi_phi_quy_{year}_Q{quarter}.xlsx",
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }

    @api.model
    def get_cost_estimate_board(self, params=None):
        """Dự toán chi phí T1→T12 — cùng nguồn Danh sách thanh toán từng tháng.

        Mỗi tháng = search_month_cost_ky(year, month) + số tiền kỳ
        (_month_ky_period_amount), giống khi lọc tháng trên Danh sách TT.
        """
        params = params or {}
        today = fields.Date.context_today(self)
        year = int(params.get("year") or today.year)
        region = (params.get("region") or "all") or "all"

        Service = self.env["phan.he.service"]
        extra = []
        if region and region != "all":
            extra.append(("store_mien", "=", region))

        month_amounts = [0.0] * 12
        month_counts = [0] * 12
        store_map = {}
        all_regions = set()

        for m in range(1, 13):
            recs = Service.search_month_cost_ky(
                year, m, extra_domain=extra or None
            )
            month_counts[m - 1] = len(recs)
            for rec in recs:
                amt = float(Service._month_ky_period_amount(rec) or 0.0)
                if amt <= 0:
                    continue
                store = rec.store_id
                # Key theo HĐ để khớp từng dòng Danh sách thanh toán
                sid = rec.id
                if sid not in store_map:
                    store_map[sid] = {
                        "id": rec.id,
                        "store": store.name or rec.name or "—",
                        "code": rec.customer_code or rec.code or "",
                        "provider": (rec.provider_id.name if rec.provider_id else "") or "—",
                        "region": rec.store_mien or "",
                        "bandwidth": rec.bandwidth or "—",
                        "monthly": 0.0,
                        "amounts": [0.0] * 12,
                        "total": 0.0,
                    }
                row = store_map[sid]
                row["amounts"][m - 1] += amt
                row["total"] += amt
                if amt > row["monthly"]:
                    row["monthly"] = amt
                month_amounts[m - 1] += amt
                if row["region"]:
                    all_regions.add(row["region"])

        months = []
        for m in range(1, 13):
            months.append({
                "month": m,
                "label": f"Tháng {m}",
                "short": f"T{m}",
                "count": int(month_counts[m - 1]),
                "amount": float(month_amounts[m - 1]),
            })

        stores = sorted(
            store_map.values(),
            key=lambda r: (r.get("region") or "", r.get("store") or ""),
        )
        for i, row in enumerate(stores, 1):
            row["stt"] = i
            row["amounts"] = [float(x) for x in row["amounts"]]
            row["total"] = float(row["total"])
            row["monthly"] = float(row["monthly"])

        year_total = float(sum(month_amounts))
        peak = max(months, key=lambda x: x["amount"]) if months else None

        # KPI Số cửa hàng = Internet Đang sử dụng (cùng badge sidebar)
        active_domain = [
            ("active", "=", True),
            ("service_type_id.code", "=", "internet"),
            ("state", "=", "active"),
        ]
        if region and region != "all":
            active_domain.append(("store_mien", "=", region))
        active_count = Service.search_count(active_domain)

        return {
            "year": year,
            "months": months,
            "stores": stores,
            "regions": sorted(all_regions),
            "kpi": {
                "year_total": year_total,
                "avg_month": year_total / 12.0 if year_total else 0.0,
                "store_count": int(active_count),
                "store_count_unique": len(stores),
                "peak_month": peak["month"] if peak and peak["amount"] else 0,
                "peak_amount": float(peak["amount"]) if peak else 0.0,
                "peak_label": peak["label"] if peak and peak["amount"] else "—",
            },
        }

    @api.model
    def export_cost_estimate_excel(self, params=None):
        """Xuất Excel dự toán chi phí T1–T12 (tóm tắt tháng + ma trận cửa hàng)."""
        import base64
        import io

        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            from odoo.exceptions import UserError

            raise UserError("Thiếu thư viện openpyxl trên server.") from exc

        params = params or {}
        data = self.get_cost_estimate_board(params)
        year = data.get("year")
        months = data.get("months") or []
        stores = data.get("stores") or []
        kpi = data.get("kpi") or {}

        search = (params.get("search") or "").strip().lower()
        table_region = params.get("table_region") or "all"
        if search or (table_region not in ("all", "", None)):
            filtered = []
            for row in stores:
                if table_region not in ("all", "", None) and row.get("region") != table_region:
                    continue
                if search:
                    blob = " ".join([
                        str(row.get("store") or ""),
                        str(row.get("code") or ""),
                        str(row.get("provider") or ""),
                        str(row.get("region") or ""),
                    ]).lower()
                    if search not in blob:
                        continue
                filtered.append(row)
            stores = filtered

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Du toan {year}"[:31]

        font_title = Font(name="Times New Roman", size=16, bold=True)
        font_header = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
        font_cell = Font(name="Times New Roman", size=10)
        font_bold = Font(name="Times New Roman", size=10, bold=True)
        fill_header = PatternFill("solid", fgColor="059669")
        fill_total = PatternFill("solid", fgColor="D1FAE5")
        fill_peach = PatternFill("solid", fgColor="FBE4D5")
        thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        money_fmt = '#,##0'
        align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_l = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_r = Alignment(horizontal="right", vertical="center")

        ws.merge_cells("A1:D1")
        ws["A1"].value = f"DỰ TOÁN CHI PHÍ INTERNET NĂM {year}"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws["A2"].value = (
            f"Tổng năm: {int(kpi.get('year_total') or 0):,} đ  |  "
            f"TB/tháng: {int(kpi.get('avg_month') or 0):,} đ  |  "
            f"{int(kpi.get('store_count') or 0)} cửa hàng"
        ).replace(",", ".")
        ws.merge_cells("A2:D2")

        # --- Bảng tóm tắt T1–T12 ---
        ws["A4"].value = "TÓM TẮT THEO THÁNG"
        ws["A4"].font = font_bold
        sum_headers = ["STT", "Tháng", "Số cửa hàng", "Tổng dự toán (đ)"]
        for col, h in enumerate(sum_headers, 1):
            c = ws.cell(5, col, h)
            c.font = font_header
            c.fill = fill_header
            c.border = thin
            c.alignment = align_c

        for idx, mo in enumerate(months, 1):
            r = 5 + idx
            vals = [idx, mo.get("label"), int(mo.get("count") or 0), float(mo.get("amount") or 0)]
            for col, val in enumerate(vals, 1):
                cell = ws.cell(r, col, val)
                cell.font = font_cell
                cell.border = thin
                cell.alignment = align_c if col <= 3 else align_r
                if col == 4:
                    cell.number_format = money_fmt

        foot = 5 + len(months) + 1
        for col, val in enumerate(
            ["", "TỔNG NĂM", int(kpi.get("store_count") or 0), float(kpi.get("year_total") or 0)],
            1,
        ):
            cell = ws.cell(foot, col, val)
            cell.font = font_bold
            cell.fill = fill_total
            cell.border = thin
            if col == 4:
                cell.number_format = money_fmt
                cell.alignment = align_r

        # --- Sheet ma trận cửa hàng × tháng ---
        ws2 = wb.create_sheet(f"Cua hang {year}"[:31])
        headers2 = ["STT", "Cửa hàng", "Mã KH", "Nhà cung cấp", "Miền"] + [
            f"T{m}" for m in range(1, 13)
        ] + ["Tổng năm"]
        ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers2))
        ws2["A1"].value = f"DỰ TOÁN CHI PHÍ THEO CỬA HÀNG — NĂM {year}"
        ws2["A1"].font = font_title
        ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 28

        for col, h in enumerate(headers2, 1):
            c = ws2.cell(3, col, h)
            c.font = font_header
            c.fill = fill_header
            c.border = thin
            c.alignment = align_c

        for idx, row in enumerate(stores, 1):
            ridx = 3 + idx
            values = [
                idx,
                row.get("store") or "",
                row.get("code") or "",
                row.get("provider") or "",
                row.get("region") or "",
            ] + [float(x) for x in (row.get("amounts") or [0.0] * 12)] + [float(row.get("total") or 0)]
            for col, val in enumerate(values, 1):
                cell = ws2.cell(ridx, col, val)
                cell.font = font_cell
                cell.border = thin
                if col <= 5:
                    cell.alignment = align_c if col == 1 else align_l
                else:
                    cell.alignment = align_r
                    cell.number_format = money_fmt

        tot_row = 4 + len(stores)
        month_totals = [0.0] * 12
        year_sum = 0.0
        for row in stores:
            amts = row.get("amounts") or [0.0] * 12
            for i in range(12):
                month_totals[i] += float(amts[i] or 0)
            year_sum += float(row.get("total") or 0)
        foot_vals = ["", "TỔNG CỘNG", "", "", ""] + month_totals + [year_sum]
        for col, val in enumerate(foot_vals, 1):
            cell = ws2.cell(tot_row, col, val)
            cell.font = font_bold
            cell.fill = fill_peach
            cell.border = thin
            if col >= 6:
                cell.number_format = money_fmt
                cell.alignment = align_r

        for sheet in (ws, ws2):
            max_col = sheet.max_column or 1
            for col in range(1, max_col + 1):
                maxlen = 8
                for r in range(1, min(sheet.max_row or 1, 80) + 1):
                    maxlen = max(maxlen, len(str(sheet.cell(r, col).value or "")))
                sheet.column_dimensions[get_column_letter(col)].width = min(28, max(10, maxlen + 2))

        buf = io.BytesIO()
        wb.save(buf)
        return {
            "file_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
            "filename": f"Du_toan_chi_phi_{year}.xlsx",
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }

